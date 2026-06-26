#!/usr/bin/env python3
"""Build a live ERP reconciliation report from MariaDB on the print server.

The script intentionally imports /var/www/print/reconciliation_api_server.py
when it runs on the server. That keeps the DB connection path identical to the
running reconciliation API and avoids copying credentials into this repository.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import math
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


DEFAULT_API_PATH = Path("/var/www/print/reconciliation_api_server.py")
DEFAULT_SPEC_NUMS = ["1063", "1064", "1065", "1068", "1072", "1073", "1074", "1076"]


@dataclass(frozen=True)
class ResolvedSpec:
    spec_id: int
    spec_num: str
    spec_num_short: str
    spec_type: str
    spec_subtype: str
    spec_subtype_id: int
    spec_date: str
    spec_name: str
    dog_id: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run live MariaDB checks for ERP/1C reconciliation and export JSON/CSV/XLSX."
    )
    parser.add_argument("--api-path", default=str(DEFAULT_API_PATH), help="Path to reconciliation_api_server.py")
    parser.add_argument("--out-dir", default="/var/www/print/akt_sverki/live_validation")
    parser.add_argument("--dog-id", type=int, default=88)
    parser.add_argument("--dog-name", default="660/1")
    parser.add_argument("--contact-id", type=int, default=115)
    parser.add_argument("--legal-id", type=int, default=221)
    parser.add_argument(
        "--spec-nums",
        default=",".join(DEFAULT_SPEC_NUMS),
        help="Comma-separated veda_specs.f_num values. Default is the current AERO-TRADE control set.",
    )
    parser.add_argument("--label", default="aero_trade_660_1_live")
    return parser.parse_args()


def load_api(api_path: Path) -> Any:
    if not api_path.exists():
        raise FileNotFoundError(f"API file not found: {api_path}")
    spec = importlib.util.spec_from_file_location("server_reconciliation_api", api_path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import API from {api_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def q(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"


def nz(value: Any) -> float:
    try:
        number = float(value or 0)
    except Exception:
        return 0.0
    if math.isfinite(number):
        return number
    return 0.0


def money(value: Any) -> float:
    return round(nz(value), 2)


def one_line(value: Any) -> str:
    return " ".join(str(value or "").replace("\t", " ").replace("\r", " ").replace("\n", " ").split())


def joined_lines(values: list[str]) -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = one_line(value)
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return "\n".join(result)


def run_sql(api: Any, sql: str) -> list[list[str]]:
    return api.run_mysql_tsv(sql)


def resolve_specs(api: Any, dog_id: int, spec_nums: list[str]) -> list[ResolvedSpec]:
    if not spec_nums:
        raise ValueError("At least one spec number is required")
    in_list = ",".join(q(item) for item in spec_nums)
    order_expr = ",".join(q(item) for item in spec_nums)
    rows = run_sql(
        api,
        f"""
SELECT
    s.f_id,
    COALESCE(s.f_num, '') AS spec_num,
    COALESCE(s.f_num, '') AS spec_num_short,
    COALESCE(NULLIF(spec_type.f_dopprstr, ''), NULLIF(spec_type.f_name, ''), '') AS spec_type,
    COALESCE(NULLIF(spec_subtype.f_dopprstr, ''), NULLIF(spec_subtype.f_name, ''), '') AS spec_subtype,
    COALESCE(s.f_subtype, 0) AS spec_subtype_id,
    COALESCE(DATE_FORMAT(s.f_dt, '%Y-%m-%d'), '') AS spec_date,
    COALESCE(s.f_tovar, '') AS spec_name,
    COALESCE(s.f_dogid, 0) AS dog_id
FROM veda_specs s
LEFT JOIN veda_spr spec_type
       ON spec_type.f_type = 33
      AND spec_type.f_num = s.f_typez
LEFT JOIN veda_spr spec_subtype
       ON spec_subtype.f_type = 130
      AND spec_subtype.f_num = s.f_subtype
WHERE s.f_dogid = {dog_id}
  AND CAST(s.f_num AS CHAR) IN ({in_list})
ORDER BY FIELD(CAST(s.f_num AS CHAR), {order_expr}), s.f_id;
""",
    )
    resolved = [
        ResolvedSpec(
            spec_id=int(row[0]),
            spec_num=row[1],
            spec_num_short=row[2],
            spec_type=row[3],
            spec_subtype=row[4],
            spec_subtype_id=int(row[5] or 0),
            spec_date=row[6],
            spec_name=row[7],
            dog_id=int(row[8] or 0),
        )
        for row in rows
    ]
    found = {item.spec_num for item in resolved}
    missing = [item for item in spec_nums if item not in found]
    if missing:
        raise RuntimeError(f"Spec numbers not found for dog_id={dog_id}: {', '.join(missing)}")
    return resolved


def fetch_context(api: Any, contact_id: int, legal_id: int, dog_id: int, dog_name: str) -> dict[str, Any]:
    rows = run_sql(
        api,
        f"""
SELECT
    c.f_id AS contact_id,
    COALESCE(NULLIF(c.f_cname, ''), NULLIF(c.f_name, ''), '') AS contact_name,
    COALESCE(c.f_inn, '') AS contact_inn,
    cl.f_id AS legal_id,
    COALESCE(cl.f_cname, '') AS legal_name,
    COALESCE(cl.f_inn, '') AS legal_inn,
    d.f_id AS dog_id,
    COALESCE(d.f_dogname, '') AS dog_name,
    COALESCE(DATE_FORMAT(d.f_dogdate, '%Y-%m-%d'), '') AS dog_date,
    COALESCE(d.f_kod1c, '') AS dog_code1c
FROM veda_contacts c
JOIN veda_clients cl ON cl.f_contactid = c.f_id
JOIN veda_dogs d ON d.f_contrid = cl.f_id
WHERE c.f_id = {contact_id}
  AND cl.f_id = {legal_id}
  AND d.f_id = {dog_id}
  AND d.f_dogname = {q(dog_name)}
LIMIT 1;
""",
    )
    if not rows:
        raise RuntimeError(
            f"Context not found: contact_id={contact_id}, legal_id={legal_id}, dog_id={dog_id}, dog_name={dog_name}"
        )
    row = rows[0]
    return {
        "contact_id": int(row[0]),
        "contact_name": row[1],
        "contact_inn": row[2],
        "legal_id": int(row[3]),
        "legal_name": row[4],
        "legal_inn": row[5],
        "dog_id": int(row[6]),
        "dog_name": row[7],
        "dog_date": row[8],
        "dog_code1c": row[9],
    }


def operation_bucket(operation: dict[str, Any]) -> str:
    marker = int(operation.get("reimbursement_id") or 0)
    if marker == 1:
        return "reimbursable"
    if marker == 2:
        return "non_reimbursable"
    return "unclassified"


def customer_invoices(erp_docs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    invoices: list[dict[str, Any]] = []
    seen: set[int] = set()
    for doc in erp_docs:
        if doc.get("doc_kind") != "schet":
            continue
        if int(doc.get("invoice_id") or 0) != 0:
            continue
        if one_line(doc.get("type_name")) != "Счет покупателю":
            continue
        doc_id = int(doc.get("erp_doc_id") or 0)
        if doc_id <= 0 or doc_id in seen:
            continue
        seen.add(doc_id)
        invoices.append(doc)
    invoices.sort(key=lambda item: (str(item.get("date_iso") or ""), str(item.get("number") or ""), int(item.get("erp_doc_id") or 0)))
    return invoices


def closing_docs(erp_docs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    docs: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, float]] = set()
    for doc in erp_docs:
        if doc.get("doc_kind") != "act":
            continue
        key = (
            one_line(doc.get("code1c")),
            one_line(doc.get("number")),
            one_line(doc.get("date")),
            money(doc.get("sum")),
        )
        if key in seen:
            continue
        seen.add(key)
        docs.append(doc)
    docs.sort(key=lambda item: (str(item.get("date_iso") or ""), str(item.get("code1c") or ""), str(item.get("number") or "")))
    return docs


def format_invoice_amount(doc: dict[str, Any]) -> str:
    amount = money(doc.get("sum"))
    currency = one_line(doc.get("currency"))
    if currency:
        return f"{amount:,.2f} {currency}".replace(",", " ")
    return f"{amount:,.2f}".replace(",", " ")


def format_doc_line(doc: dict[str, Any]) -> str:
    number = one_line(doc.get("code1c")) or one_line(doc.get("number"))
    date = one_line(doc.get("date"))
    amount = money(doc.get("sum"))
    if number and date:
        return f"{number} от {date} ({amount:,.2f} р.)".replace(",", " ")
    if number:
        return f"{number} ({amount:,.2f} р.)".replace(",", " ")
    return ""


def build_spec_block(api: Any, item: ResolvedSpec) -> dict[str, Any]:
    delivery = api.fetch_delivery(item.spec_id)
    operations = api.fetch_operations(item.spec_id)
    erp_docs = api.fetch_erp_docs(item.spec_id)
    payments = api.fetch_payments(item.spec_id)
    invoices = customer_invoices(erp_docs)
    acts = closing_docs(erp_docs)

    paid_by_routine = money(sum(nz(op.get("rp_paid_sum")) for op in operations))
    paid_by_docs = money(sum(nz(pay.get("classified_sum")) for pay in payments if pay.get("direction") == "incoming"))
    reimbursable = money(sum(nz(op.get("rp_realiz_sum")) for op in operations if operation_bucket(op) == "reimbursable"))
    non_reimbursable = money(sum(nz(op.get("rp_realiz_sum")) for op in operations if operation_bucket(op) == "non_reimbursable"))
    unclassified_realization = money(sum(nz(op.get("rp_realiz_sum")) for op in operations if operation_bucket(op) == "unclassified"))
    realization_total = money(reimbursable + non_reimbursable + unclassified_realization)
    delta = money(paid_by_routine - realization_total)
    invoice_total = money(sum(nz(doc.get("sum")) for doc in invoices))

    issues: list[str] = []
    if not invoices:
        issues.append("NO_CUSTOMER_INVOICE")
    if not acts:
        issues.append("NO_CLOSING_DOC")
    if abs(paid_by_routine - paid_by_docs) > 0.01:
        issues.append("PAYMENT_ROUTINE_VS_DOCS_MISMATCH")
    if abs(delta) > 0.01:
        issues.append("SETTLEMENT_BALANCE_NONZERO")
    if unclassified_realization:
        issues.append("UNCLASSIFIED_REALIZATION")

    operation_rows = []
    for op in operations:
        operation_rows.append(
            {
                "oper_id": op.get("oper_id"),
                "oper_num": op.get("oper_num"),
                "type": op.get("oper_type_name"),
                "f_isvozm": op.get("reimbursement_id"),
                "isvozm_name": op.get("reimbursement_name"),
                "nds": op.get("nds_name"),
                "paid_get_paidsum": money(op.get("rp_paid_sum")),
                "realiz_get_realizsum": money(op.get("rp_realiz_sum")),
                "expenses_get_expensessum": money(op.get("rp_expenses_sum")),
                "profit_get_profit": money(op.get("rp_profit_sum")),
            }
        )

    return {
        "spec_id": item.spec_id,
        "spec_num": item.spec_num,
        "spec_num_short": item.spec_num_short,
        "spec_type": item.spec_type,
        "spec_subtype": item.spec_subtype,
        "spec_subtype_id": item.spec_subtype_id,
        "spec_label": f"{item.spec_type or 'Поставка'} №{item.spec_num}",
        "spec_date": item.spec_date,
        "spec_name": item.spec_name,
        "delivery": delivery,
        "totals": {
            "customer_invoice_total": invoice_total,
            "paid_total_get_paidsum": paid_by_routine,
            "paid_total_acchist_docs": paid_by_docs,
            "reimbursable_realization_get_realizsum": reimbursable,
            "non_reimbursable_realization_get_realizsum": non_reimbursable,
            "unclassified_realization_get_realizsum": unclassified_realization,
            "realization_total_get_realizsum": realization_total,
            "settlement_delta_paid_minus_realization": delta,
        },
        "counts": {
            "operations": len(operations),
            "customer_invoices": len(invoices),
            "closing_docs": len(acts),
            "payments": len(payments),
        },
        "issues": issues,
        "customer_invoices": [
            {
                "erp_doc_id": doc.get("erp_doc_id"),
                "number": doc.get("number"),
                "date": doc.get("date"),
                "sum": money(doc.get("sum")),
                "currency": doc.get("currency"),
                "code1c": doc.get("code1c"),
            }
            for doc in invoices
        ],
        "closing_docs": [
            {
                "erp_doc_id": doc.get("erp_doc_id"),
                "number": doc.get("number"),
                "date": doc.get("date"),
                "sum": money(doc.get("sum")),
                "currency": doc.get("currency"),
                "code1c": doc.get("code1c"),
                "type_name": doc.get("type_name"),
                "nds": doc.get("nds_name"),
            }
            for doc in acts
        ],
        "operation_rows": operation_rows,
    }


def build_report(api: Any, args: argparse.Namespace) -> dict[str, Any]:
    spec_nums = [item.strip() for item in str(args.spec_nums).split(",") if item.strip()]
    context = fetch_context(api, args.contact_id, args.legal_id, args.dog_id, args.dog_name)
    specs = resolve_specs(api, args.dog_id, spec_nums)

    blocks: list[dict[str, Any]] = []
    for index, spec_item in enumerate(specs, start=1):
        print(f"[{index}/{len(specs)}] spec {spec_item.spec_num} ({spec_item.spec_id})", file=sys.stderr)
        blocks.append(build_spec_block(api, spec_item))

    summary = {
        "specs": len(blocks),
        "customer_invoice_total": money(sum(block["totals"]["customer_invoice_total"] for block in blocks)),
        "paid_total_get_paidsum": money(sum(block["totals"]["paid_total_get_paidsum"] for block in blocks)),
        "paid_total_acchist_docs": money(sum(block["totals"]["paid_total_acchist_docs"] for block in blocks)),
        "reimbursable_realization_get_realizsum": money(
            sum(block["totals"]["reimbursable_realization_get_realizsum"] for block in blocks)
        ),
        "non_reimbursable_realization_get_realizsum": money(
            sum(block["totals"]["non_reimbursable_realization_get_realizsum"] for block in blocks)
        ),
        "realization_total_get_realizsum": money(sum(block["totals"]["realization_total_get_realizsum"] for block in blocks)),
        "settlement_delta_paid_minus_realization": money(
            sum(block["totals"]["settlement_delta_paid_minus_realization"] for block in blocks)
        ),
        "issues": {},
    }
    issue_counts: dict[str, int] = {}
    for block in blocks:
        for issue in block["issues"]:
            issue_counts[issue] = issue_counts.get(issue, 0) + 1
    summary["issues"] = issue_counts

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source": {
            "api_path": str(args.api_path),
            "db_host": getattr(api, "DB_HOST", ""),
            "db_name": getattr(api, "DB_NAME", ""),
            "db_user": getattr(api, "DB_USER", ""),
            "connection_note": "DB secret is read only by the server API module and is not exported.",
        },
        "parameters": {
            "contact_id": args.contact_id,
            "legal_id": args.legal_id,
            "dog_id": args.dog_id,
            "dog_name": args.dog_name,
            "spec_nums": spec_nums,
        },
        "context": context,
        "sql_sources": {
            "hierarchy": "veda_contacts -> veda_clients.f_contactid -> veda_dogs.f_contrid -> veda_specs.f_dogid; label/type from veda_spr(f_type=33/130), view_specinv/view_specs are not production sources",
            "operations": "veda_spec_invoices with f_parenttype in (2,4); f_parenttype=4 resolved through veda_categs f_ctgtype=24/f_objecttype=5",
            "payments": "get_paidsum(oper.f_id) and control SUM(veda_acchist_docs.f_clssum); JOIN veda_acchist ah ON ah.f_id = ahd.f_acchistid; ahd.f_doctype=3; ah.f_type=0",
            "realization": "get_realizsum(oper.f_id), split by veda_spec_invoices.f_isvozm",
            "customer_invoices": "main customer invoices selected from server API docs as type_name='Счет покупателю' and invoice_id=0; equivalent to main spec invoices, not supplier invoice rows",
            "closing_docs": "veda_akts through server API fetch_erp_docs; every document is listed with a line break in XLSX",
            "delta": "(+/-) = SUM(get_paidsum) - SUM(get_realizsum)",
        },
        "summary": summary,
        "blocks": blocks,
    }


def write_json(report: dict[str, Any], path: Path) -> None:
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(report: dict[str, Any], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(
            [
                "spec_num",
                "spec_label",
                "spec_type",
                "spec_subtype",
                "spec_id",
                "invoice_total",
                "paid_get_paidsum",
                "paid_acchist_docs",
                "reimbursable_get_realizsum",
                "non_reimbursable_get_realizsum",
                "realization_total_get_realizsum",
                "delta_paid_minus_realization",
                "customer_invoices",
                "closing_docs",
                "issues",
            ]
        )
        for block in report["blocks"]:
            writer.writerow(
                [
                    block["spec_num"],
                    block.get("spec_label", ""),
                    block.get("spec_type", ""),
                    block.get("spec_subtype", ""),
                    block["spec_id"],
                    block["totals"]["customer_invoice_total"],
                    block["totals"]["paid_total_get_paidsum"],
                    block["totals"]["paid_total_acchist_docs"],
                    block["totals"]["reimbursable_realization_get_realizsum"],
                    block["totals"]["non_reimbursable_realization_get_realizsum"],
                    block["totals"]["realization_total_get_realizsum"],
                    block["totals"]["settlement_delta_paid_minus_realization"],
                    joined_lines([str(inv["number"]) for inv in block["customer_invoices"]]),
                    joined_lines([format_doc_line(doc) for doc in block["closing_docs"]]),
                    ", ".join(block["issues"]),
                ]
            )


def write_xlsx(report: dict[str, Any], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Выгрузка"
    headers = [
        "№ спецификации",
        "",
        "Счет",
        "Сумма по счету",
        "Сумма оплаты",
        "",
        "Возмещаемые расходы",
        "Невозмещаемые расходы",
        "",
        "№ счф",
        "(+/-)",
    ]
    ws.append(headers)

    fills = {
        "header": PatternFill("solid", fgColor="D9D7D2"),
        "spec": PatternFill("solid", fgColor="F3F1ED"),
        "invoice": PatternFill("solid", fgColor="EEF6FF"),
        "payment": PatternFill("solid", fgColor="EEF8F0"),
        "expense": PatternFill("solid", fgColor="FFF5DF"),
        "sf": PatternFill("solid", fgColor="F7F4EF"),
        "delta_ok": PatternFill("solid", fgColor="E2F0D9"),
        "delta_bad": PatternFill("solid", fgColor="FCE4D6"),
        "sep": PatternFill("solid", fgColor="FAFAF9"),
    }
    border = Border(
        left=Side(style="thin", color="D6D3CD"),
        right=Side(style="thin", color="D6D3CD"),
        top=Side(style="thin", color="D6D3CD"),
        bottom=Side(style="thin", color="D6D3CD"),
    )
    for cell in ws[1]:
        cell.fill = fills["header"]
        cell.font = Font(bold=True, color="1F2933")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    row = 2
    merge_columns = [1, 5, 7, 8, 10, 11]
    for block in report["blocks"]:
        invoices = block["customer_invoices"] or [{"number": "", "sum": "", "currency": ""}]
        start = row
        for invoice in invoices:
            ws.cell(row=row, column=3, value=invoice.get("number") or "")
            ws.cell(row=row, column=4, value=format_invoice_amount(invoice))
            row += 1
        end = row - 1

        ws.cell(start, 1, block.get("spec_label") or f"Поставка №{block['spec_num']}")
        ws.cell(start, 5, block["totals"]["paid_total_get_paidsum"])
        ws.cell(start, 7, block["totals"]["reimbursable_realization_get_realizsum"])
        ws.cell(start, 8, block["totals"]["non_reimbursable_realization_get_realizsum"])
        ws.cell(start, 10, joined_lines([format_doc_line(doc) for doc in block["closing_docs"]]))
        ws.cell(start, 11, block["totals"]["settlement_delta_paid_minus_realization"])

        if end > start:
            for col in merge_columns:
                ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)

        for r in range(start, end + 1):
            for c in range(1, 12):
                cell = ws.cell(r, c)
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(r, 1).fill = fills["spec"]
            ws.cell(r, 2).fill = fills["sep"]
            ws.cell(r, 3).fill = fills["invoice"]
            ws.cell(r, 4).fill = fills["invoice"]
            ws.cell(r, 5).fill = fills["payment"]
            ws.cell(r, 6).fill = fills["sep"]
            ws.cell(r, 7).fill = fills["expense"]
            ws.cell(r, 8).fill = fills["expense"]
            ws.cell(r, 9).fill = fills["sep"]
            ws.cell(r, 10).fill = fills["sf"]
            ws.cell(r, 11).fill = fills["delta_ok"] if abs(block["totals"]["settlement_delta_paid_minus_realization"]) <= 0.01 else fills["delta_bad"]
        for col in [5, 7, 8, 11]:
            ws.cell(start, col).number_format = '#,##0.00" р.";[Red]-#,##0.00" р.";0.00" р."'
            ws.cell(start, col).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        ws.cell(start, 1).font = Font(bold=True)
        ws.row_dimensions[start].height = max(28, 16 * max(1, len(str(ws.cell(start, 10).value or "").splitlines())))

    total_row = row + 1
    ws.cell(total_row, 1, "ИТОГО")
    ws.cell(total_row, 5, report["summary"]["paid_total_get_paidsum"])
    ws.cell(total_row, 7, report["summary"]["reimbursable_realization_get_realizsum"])
    ws.cell(total_row, 8, report["summary"]["non_reimbursable_realization_get_realizsum"])
    ws.cell(total_row, 11, report["summary"]["settlement_delta_paid_minus_realization"])
    for c in range(1, 12):
        cell = ws.cell(total_row, c)
        cell.fill = fills["header"]
        cell.border = border
        cell.font = Font(bold=True)
    for col in [5, 7, 8, 11]:
        ws.cell(total_row, col).number_format = '#,##0.00" р.";[Red]-#,##0.00" р.";0.00" р."'

    rules = wb.create_sheet("SQL и правила")
    rules_rows = [
        ("Поле", "Источник/SQL", "Правило"),
        ("Иерархия", report["sql_sources"]["hierarchy"], "клиент -> ЮЛ -> договор -> поставка"),
        ("Операции", report["sql_sources"]["operations"], "основной источник строк поставки"),
        ("Сумма оплаты", report["sql_sources"]["payments"], "контроль get_paidsum и veda_acchist_docs.f_clssum"),
        ("Возмещаемые/невозмещаемые", report["sql_sources"]["realization"], "деление по f_isvozm, сумма из get_realizsum"),
        ("(+/-)", report["sql_sources"]["delta"], "сальдо взаиморасчетов"),
        ("СФ/УПД", report["sql_sources"]["closing_docs"], "каждый документ отдельной строкой внутри объединенной ячейки"),
    ]
    for row_values in rules_rows:
        rules.append(row_values)
    for cell in rules[1]:
        cell.fill = fills["header"]
        cell.font = Font(bold=True)
    for row_cells in rules.iter_rows():
        for cell in row_cells:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 22,
        "B": 3,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 3,
        "G": 20,
        "H": 22,
        "I": 3,
        "J": 48,
        "K": 18,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for sheet in [rules]:
        for col_idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 22 if col_idx > 1 else 16
    ws.freeze_panes = "A2"
    rules.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    args = parse_args()
    api = load_api(Path(args.api_path))
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    report = build_report(api, args)

    base = out_dir / args.label
    json_path = base.with_suffix(".json")
    csv_path = base.with_suffix(".csv")
    xlsx_path = base.with_suffix(".xlsx")
    write_json(report, json_path)
    write_csv(report, csv_path)
    write_xlsx(report, xlsx_path)

    print(json.dumps({"json": str(json_path), "csv": str(csv_path), "xlsx": str(xlsx_path), "summary": report["summary"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
