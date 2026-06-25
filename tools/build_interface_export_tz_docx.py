from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix.docx"
OUT_VERSIONED = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix_LIVE_MARIADB_2026-06-24.docx"
PDF = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix.pdf"
PDF_VERSIONED = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix_LIVE_MARIADB_2026-06-24.pdf"

IMG_UI = ROOT / "screenshots" / "ui" / "matrix_current_view.png"
IMG_XLSX = ROOT / "screenshots" / "excel" / "aero_trade_660_1_live_export.png"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(34, 34, 34)
MUTED = RGBColor(92, 92, 92)
RED = RGBColor(155, 28, 28)
GREEN = RGBColor(42, 125, 72)
FILL_HEADER = "E8EEF5"
FILL_NOTE = "F4F6F9"
FILL_WARN = "FFF1EB"


def font(run, size=10, bold=False, color=DARK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width(cell, widths[idx])
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def cell_text(cell, text, size=8.2, bold=False, color=DARK, fill=None, align=None):
    cell.text = ""
    if fill:
        shade(cell, fill)
    margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, size=8.0, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell_text(table.rows[0].cells[idx], header, size=size, bold=True, color=BLUE, fill=FILL_HEADER, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            align = aligns[idx] if aligns and idx < len(aligns) else None
            cell_text(cells[idx], value, size=size, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def para(doc, text, size=10.2, color=DARK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    font(run, size=size, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    font(run, size=9.7)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.text = ""
    run = p.add_run(text)
    font(run, size={1: 15.5, 2: 12.5, 3: 11.2}.get(level, 10), bold=True, color=BLUE if level < 3 else DARK)


def note(doc, title, text, fill=FILL_NOTE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table_geometry(table, [9360])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    font(r, size=9.8, bold=True, color=BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    font(r2, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.02
    for idx, line in enumerate(text.strip().splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        font(r, size=7.25, name="Courier New", color=RGBColor(40, 40, 40))


def image(doc, path, caption, img_width=6.45):
    if not path.exists():
        note(doc, "Скриншот не найден", str(path), FILL_WARN)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(img_width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    font(r, size=8.1, color=MUTED)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(9.8)
        style.paragraph_format.left_indent = Inches(0.32)
        style.paragraph_format.first_line_indent = Inches(-0.16)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = ""
    r = footer.add_run("Матрица акта сверки ERP ↔ 1С")
    font(r, size=8, color=MUTED)


def title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ТЗ: матрица акта сверки ERP ↔ 1С")
    font(r, size=20, bold=True, color=RGBColor(0, 0, 0))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Единый файл требований: интерфейс и XLSX-выгрузка")
    font(r, size=12.5, color=MUTED)
    add_table(
        doc,
        ["Параметр", "Требование"],
        [
            ("Контур", "Встроенный экран текущей PHP ERP, MariaDB и существующий SOAP-обмен с 1С."),
            ("Источник ERP", "Только исходные таблицы ERP и существующие процедуры. `view_specinv/view_specs` не использовать как production-источник связей."),
            ("Источник 1С", "Существующий SOAP-сервис 1С: контрагенты, договоры, счета, выписки, акты, ЭДО/СФ/УПД, акты сверки."),
            ("Разрез сверки", "Клиент -> ЮЛ -> договор -> поставка/заявка -> документы внутри поставки."),
            ("Формула дельты", "`оплаты - возмещаемые расходы - невозмещаемые расходы`. Положительно = переплата, отрицательно = долг."),
            ("Ограничение запуска", "Не запускать без клиента и периода/договора. Экранная выборка до 50 поставок, большие объемы - пагинация или фоновая XLSX-задача."),
        ],
        [1850, 7510],
        size=8.4,
    )


def block_interface(doc):
    heading(doc, "Блок 1. Интерфейс матрицы", 1)
    para(
        doc,
        "Экран нужен бухгалтерии для быстрой сверки поставок по клиенту/ЮЛ/договору. На верхних уровнях показываются только агрегаты; документы раскрываются внутри поставки.",
    )
    image(doc, IMG_UI, "Скриншот 1. Матрица по client_id=221, dog_id=88: закрепленная иерархия, итоговые суммы, `Переплата`, результат сверки 1С.", 6.45)

    heading(doc, "1.1 Запуск и фильтры", 2)
    add_table(
        doc,
        ["Фильтр / URL", "Значение для реализации"],
        [
            ("`client_id + scope=legal`", "Открывает одно ЮЛ из `veda_clients`. Пример: `client_id=221&scope=legal`."),
            ("`contact_id` или `scope=contact`", "Открывает клиента `veda_contacts` и все его ЮЛ из `veda_clients.f_contactid`."),
            ("`dog_id`", "Ограничивает поставки одним договором `veda_dogs.f_id`."),
            ("`limit`", "Ограничивает число поставок на экране. Рекомендованный максимум для UI - 50."),
            ("`compare_1c=1`", "Запускает batch-сверку с 1С и добавляет статусы `1С: X/Y совпало`, `Нет в 1С`, `Сумма расходится` и т.д."),
            ("Поиск", "Ищет по ЮЛ, ИНН, `contact_id`, `client_id`, договору, заявке/спецификации, счету, акту, СФ/УПД и номеру 1С."),
        ],
        [2500, 6860],
        size=7.9,
    )

    heading(doc, "1.2 Иерархия и колонки", 2)
    add_table(
        doc,
        ["Уровень / колонка", "ERP-источник", "Правило отображения"],
        [
            ("Клиент", "`veda_contacts`", "Агрегат всех ЮЛ клиента. Показывать сумму счетов, оплат, расходов, долг/переплату, количество проблем."),
            ("ЮЛ", "`veda_clients`", "Агрегат одного юридического лица. Фильтруется через `client_id`."),
            ("Договор", "`veda_dogs`", "Агрегат поставок договора. Номер и код 1С обязательны в подписи."),
            ("Поставка", "`veda_specs` + `veda_spr(f_type=33/130)`", "Показывать реальный тип: `Заявка`, `Спецификация`, `Сертификация` и т.д.; не хардкодить `Спецификация`."),
            ("Счет", "`veda_schets`, только `f_type=1`", "Выводить покупательские счета, поставщиков исключать."),
            ("Сумма оплаты", "`SUM(veda_acchist_docs.f_clssum)`", "Брать распределенную долю платежа, не полную банковскую выписку."),
            ("Возм./невозм. расходы", "`get_realizsum(oper_id)` + `veda_spec_invoices.f_isvozm`", "`f_isvozm=1` - возмещаемые, `f_isvozm=2` - невозмещаемые."),
            ("(+/-)", "оплаты - возм. - невозм.", "Если > 0, показывать термин `Переплата`; если < 0, `Долг`. Старый термин `Остаток` не использовать."),
        ],
        [1700, 3100, 4560],
        size=7.45,
    )

    heading(doc, "1.3 Статусы на экране", 2)
    add_table(
        doc,
        ["Статус", "Значение", "Что сравниваем"],
        [
            ("`1С: X/Y совпало`", "Для поставки найдено Y строк сверки ERP ↔ 1С, X строк полностью совпали.", "код/номер/дата, организация, контрагент, договор, сумма, валюта, НДС при наличии"),
            ("`Переплата`", "Оплаты больше суммы возмещаемых и невозмещаемых расходов.", "`SUM(get_paidsum) - SUM(get_realizsum) > 0`"),
            ("`Долг`", "Оплаты меньше суммы возмещаемых и невозмещаемых расходов.", "`SUM(get_paidsum) - SUM(get_realizsum) < 0`"),
            ("`Нет в 1С`", "ERP-документ есть, но 1С не вернула соответствующий документ.", "код 1С, номер, дата, организация, контрагент, договор"),
            ("`Нет в ERP`", "1С вернула документ, которому нет ERP-аналога.", "номер/дата/организация/контрагент/договор/сумма"),
            ("`Сумма расходится`", "Документ найден в обеих системах, но сумма отличается больше допуска.", "сумма счета, оплаты, акта или расхода"),
            ("`Номер/Дата/Договор расходится`", "Документ найден, но расходится конкретное поле.", "backend обязан вернуть конкретное поле; общий текст `Реквизиты расходятся` не использовать"),
            ("`Нет СФ/УПД`", "После загрузки источников нет закрывающего документа.", "наличие акта/УПД/СФ, не возмещаемость"),
            ("`Вопрос по НДС`", "Расходится ставка НДС ERP и 1С.", "только ставка/сумма НДС; `f_isvozm` не является причиной этого статуса"),
        ],
        [1850, 3680, 3830],
        size=7.35,
    )

    note(
        doc,
        "UX-правила",
        "Первая колонка `Иерархия` закреплена слева; деньги выравниваются вправо; родительские строки жирнее; документы не раскрываются массово; режим `Ошибки` фильтрует только строки с проблемами; журнал расхождений пишет тип ошибки, источник, операцию и суммы.",
    )


def block_export(doc):
    heading(doc, "Блок 2. XLSX-выгрузка", 1)
    para(
        doc,
        "XLSX не является копией экранной таблицы. Это бухгалтерская форма, где одна поставка занимает вертикальный блок строк счетов, а расходы, СФ/УПД и дельта объединяются по высоте блока.",
    )
    image(doc, IMG_XLSX, "Скриншот 2. Требуемый XLSX: объединенные ячейки, документы через Enter, итоговая дельта по поставке.", 6.45)

    heading(doc, "2.1 Поведение кнопки", 2)
    add_table(
        doc,
        ["Элемент", "Требование"],
        [
            ("Кнопка", "`Выгрузить XLSX` на экране матрицы."),
            ("Endpoint", "`/api/reconciliation/client-matrix.xlsx?client_id=...&dog_id=...&limit=...&scope=...&compare_1c=1`."),
            ("Файл", "Workbook минимум с листами `Выгрузка` и `Правила`."),
            ("Копирование", "Кнопка `Скопировать` формирует плоский TSV: значения объединенных ячеек повторяются в каждой строке."),
            ("Производительность", "Если выборка больше 50 поставок или SOAP-ответ тяжелый, формировать XLSX фоновой задачей и отдавать ссылку на готовый файл."),
        ],
        [2100, 7260],
        size=7.8,
    )

    heading(doc, "2.2 Колонки и источники", 2)
    add_table(
        doc,
        ["Колонка XLSX", "Откуда берем", "Правило"],
        [
            ("№ спецификации", "`veda_specs.f_typez/f_subtype/f_num/f_id` + `veda_spr(f_type=33/130)`", "Объединить по числу строк счетов. Тип выводить как `Заявка №1077`, `Спецификация №...`, а не хардкодить."),
            ("Счет", "`veda_schets`, основной счет, `f_type=1`", "Каждый покупательский счет отдельной строкой."),
            ("Сумма по счету", "`veda_schets.f_sum`, валюта счета", "Деньги вправо. Разные валюты не суммировать без курса."),
            ("Сумма оплаты", "`SUM(veda_acchist_docs.f_clssum)` по операциям поставки", "Объединить по поставке. Контроль: сумма должна совпадать с `SUM(get_paidsum(oper_id))`."),
            ("Возмещаемые расходы", "`SUM(get_realizsum(oper_id)) WHERE f_isvozm=1`", "Объединить по поставке."),
            ("Невозмещаемые расходы", "`SUM(get_realizsum(oper_id)) WHERE f_isvozm=2`", "Объединить по поставке."),
            ("№ счф", "`veda_akts`, СФ/УПД/акты + данные 1С/ЭДО", "Все документы выводить через Enter внутри объединенной ячейки, без `+3 документов`."),
            ("(+/-)", "`SUM(get_paidsum) - SUM(get_realizsum)`", "Объединить по поставке; положительно = переплата, отрицательно = долг."),
        ],
        [1700, 4100, 3560],
        size=7.1,
    )

    heading(doc, "2.3 Минимальные SQL-правила", 2)
    code(
        doc,
        """
-- 1. Операции поставки: прямые и f_parenttype=4
SELECT si.*
FROM veda_spec_invoices si
LEFT JOIN veda_categs cg
  ON cg.f_objectid = si.f_id AND cg.f_ctgtype = 24 AND cg.f_objecttype = 5
WHERE (si.f_parenttype = 2 AND si.f_specid = :spec_id)
   OR (si.f_parenttype = 4 AND CAST(cg.f_valstr AS SIGNED) = :spec_id);

-- 2. Покупательские счета
SELECT main.f_id, main.f_num, main.f_sum, main.f_val, main.f_kod1c, main.f_dt
FROM veda_schets child
JOIN veda_schets main ON main.f_id = IF(child.f_maininv > 0, child.f_maininv, child.f_id)
WHERE child.f_operid IN (:operation_ids)
  AND main.f_type = 1;

-- 3. Распределенная оплата
SELECT ahd.f_docid operation_id, SUM(ahd.f_clssum) paid_sum
FROM veda_acchist_docs ahd
JOIN veda_acchist ah ON ah.f_id = ahd.f_acchistid
WHERE ahd.f_doctype = 3
  AND ahd.f_docid IN (:operation_ids)
  AND ah.f_type = 0
GROUP BY ahd.f_docid;
""",
    )

    heading(doc, "2.4 Проверки приемки XLSX", 2)
    for item in [
        "Одна поставка = один вертикальный блок строк счетов.",
        "`№ спецификации`, `Сумма оплаты`, `Возмещаемые`, `Невозмещаемые`, `№ счф`, `(+/-)` объединены по высоте блока.",
        "Все акты/СФ/УПД отражены в ячейке `№ счф` через Enter, без сокращения списка.",
        "Лист `Правила` содержит краткое описание формул, источников и ограничения `view_specinv/view_specs не использовать`.",
        "Итоговая строка подписана `ИТОГО`, без номеров строк внешних таблиц.",
        "Долг/переплата считается по той же формуле, что и экран: оплаты минус реализация.",
    ]:
        bullet(doc, item)


def build():
    doc = Document()
    configure(doc)
    title(doc)
    block_interface(doc)
    block_export(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    shutil.copyfile(OUT, OUT_VERSIONED)
    for pdf in (PDF, PDF_VERSIONED):
        if pdf.exists():
            pdf.unlink()
    print(OUT)


if __name__ == "__main__":
    build()
