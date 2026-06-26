from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix.docx"
VERSION_LABEL = "PHP ERP + MariaDB + сервис 1С / 2026-06-25 / финальное ТЗ для ERP-разработчика"
OUT_VERSIONED = ROOT / "docs" / "TZ_ERP_1C_Akt_Sverki_Matrix_LIVE_MARIADB_2026-06-25.docx"

IMG_SALDO = ROOT / "screenshots" / "ui" / "saldo_live_limit3.png"
IMG_MONITOR = ROOT / "screenshots" / "ui" / "erp_1c_monitor_live_limit3.png"
IMG_MONITOR_ISSUES = ROOT / "screenshots" / "ui" / "erp_1c_monitor_live_limit3_issues.png"
IMG_XLSX = ROOT / "screenshots" / "excel" / "aero_trade_660_1_live_export.png"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
INK = RGBColor(30, 30, 30)
RISK = RGBColor(155, 28, 28)
OK_GREEN = RGBColor(42, 125, 72)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WARN_FILL = "FFF1EB"
WHITE = "FFFFFF"


def set_font(run, size=None, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, *, size=8.5, bold=False, color=INK, align=None, fill=None):
    cell.text = ""
    if fill:
        shade_cell(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    run = p.add_run("" if text is None else str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, font_size=8.2, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        set_cell_text(
            table.rows[0].cells[i],
            header,
            size=font_size,
            bold=True,
            color=DARK_BLUE,
            fill=HEADER_FILL,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            align = alignments[i] if alignments and i < len(alignments) else None
            set_cell_text(cells[i], value, size=font_size, align=align)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.text = ""
    run = p.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return p


def add_body(doc, text, *, color=INK, after=6, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, size=10.5, bold=True, color=color)
        rest = text[len(bold_prefix):]
        if rest:
            run = p.add_run(rest)
            set_font(run, size=10.5, color=color)
    else:
        run = p.add_run(text)
        set_font(run, size=10.5, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, size=10)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, size=10)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.line_spacing = 1.05
    for line_no, line in enumerate(text.strip("\n").splitlines()):
        if line_no:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, size=7.6, name="Courier New", color=RGBColor(40, 40, 40))
    return p


def add_callout(doc, title, text, fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_font(run, size=10, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    run = p2.add_run(text)
    set_font(run, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, path, caption, width=6.3):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap_run = cap.add_run(caption)
    set_font(cap_run, size=8.3, italic=True, color=MUTED)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("Техническое задание для PHP-разработчика")
    set_font(run, size=20, bold=True, color=RGBColor(0, 0, 0))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("Сальдо по поставкам и отдельная сверка ERP ↔ 1С")
    set_font(run, size=13, color=MUTED)

    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ("Версия", VERSION_LABEL),
            ("Контур", "Встроенный экран текущей PHP ERP, MariaDB, существующий SOAP-обмен с 1С."),
            ("Назначение", "Разделить две задачи: бухгалтерскую сальдовку по поставкам и технический монитор полноты ERP ↔ 1С."),
            ("Итог", "Экран сальдо по поставкам с XLSX и отдельная страница `Монитор ERP ↔ 1С` с протоколом расхождений."),
            ("Не входит", "новая отдельная система, новая БД, замена действующего ERP-интерфейса, внешние аналитические файлы как источник расчета."),
            ("Контрольный пример", "АЭРО-ТРЕЙД, договор 660/1, свежие поставки 1063, 1064, 1065, 1068, 1072, 1073, 1074, 1076."),
        ],
        [1900, 7460],
        font_size=8.8,
    )
    add_callout(
        doc,
        "Главное правило",
        "Текущая матрица не должна называться полноценной сверкой ERP ↔ 1С. Это ERP-сальдо по поставкам. Проверка соответствия двух учетных систем должна быть отдельным экраном `Монитор ERP ↔ 1С` и отдельным процессом сверки.",
    )


def add_user_requirements(doc):
    add_heading(doc, "1. Пользовательские требования", 1)
    add_body(
        doc,
        "Цель доработки делится на две пользовательские задачи. Первая - дать бухгалтеру рабочий экран сальдо по поставкам: счета, оплаты, возмещаемые и невозмещаемые расходы, долг или переплата по выбранному набору поставок. Вторая - дать отдельный монитор проверки, что документы и сальдо ERP действительно совпадают с 1С.",
    )
    add_table(
        doc,
        ["Роль", "Что пользователь хочет сделать", "Результат для пользователя"],
        [
            (
                "Бухгалтер",
                "Выбрать клиента, ЮЛ, договор и период, затем получить дерево поставок с расчетами по каждой поставке.",
                "Видит сальдо по поставкам и может выгрузить бухгалтерскую XLSX-форму.",
            ),
            (
                "Бухгалтер",
                "Запустить отдельный монитор ERP ↔ 1С по ЮЛ, договору, периоду или поставке.",
                "Получает протокол: что есть в ERP, чего нет в 1С, что есть в 1С, чего нет в ERP, где расходятся суммы, даты, номера, договоры/заявки, НДС и сальдо.",
            ),
            (
                "Бухгалтер",
                "Выгрузить сверку в привычный Excel-формат для передачи клиенту или внутренней проверки.",
                "Получает XLSX с объединенными ячейками по поставке и полным списком документов через перенос строки.",
            ),
            (
                "Руководитель бухгалтерии",
                "Видеть частотность проблем по типам операций, документам и источникам.",
                "Получает журнал расхождений отдельной сверки ERP ↔ 1С, по которому можно анализировать причины и объем проблем.",
            ),
            (
                "Разработчик ERP",
                "Реализовать два разных экрана без смешения задач.",
                "Использует описанные SQL-связи, SOAP-методы, формулы сальдо, протокол расхождений и критерии приемки.",
            ),
        ],
        [1700, 4450, 3210],
        font_size=7.9,
    )
    add_heading(doc, "1.1 Границы решения", 2)
    add_table(
        doc,
        ["Входит в scope", "Не входит в scope"],
        [
            (
                "Экран `Сальдо по поставкам` внутри текущей PHP ERP, работающий только по данным MariaDB ERP.",
                "Называть этот экран полноценной сверкой ERP ↔ 1С.",
            ),
            (
                "Отдельный экран `Сверка ERP ↔ 1С` с запуском проверки, протоколом расхождений и сравнением итогов.",
                "Использование внешних аналитических файлов как production-источника данных.",
            ),
            (
                "Сверка ERP-документов с 1С: счета покупателю, оплаты, акты/УПД/СФ, договоры и итоговые сальдо на отдельной странице.",
                "Ручное подтягивание внешних писем или документов, которых нет в ERP и 1С.",
            ),
            (
                "Журналирование кодов расхождений для последующей аналитики проблем.",
                "Изменение бизнес-логики формирования первичных документов в ERP или 1С.",
            ),
        ],
        [4680, 4680],
        font_size=7.85,
    )


def add_product_split(doc):
    add_heading(doc, "1.2 Декомпозиция на два экрана", 2)
    add_body(
        doc,
        "Чтобы не смешивать разные пользовательские цели, решение должно состоять из двух самостоятельных страниц. Первая закрывает бухгалтерское сальдо по поставкам. Вторая закрывает контроль полноты и совпадения ERP и 1С.",
    )
    add_table(
        doc,
        ["Страница", "Назначение", "Источник данных", "Что является результатом"],
        [
            (
                "Сальдо по поставкам",
                "Показать клиенту/бухгалтеру, сколько выставлено, оплачено, отражено возмещаемых/невозмещаемых расходов, какой долг или переплата по поставкам.",
                "MariaDB ERP: `veda_specs`, `veda_spec_invoices`, `veda_schets`, `veda_acchist_docs`, `get_paidsum`, `get_realizsum`.",
                "Экран с агрегатами по поставкам и XLSX в бухгалтерском формате. Это не доказательство совпадения ERP и 1С.",
            ),
            (
                "Монитор ERP ↔ 1С",
                "Проверить, что учетные системы совпадают: все документы ERP есть в 1С, все документы 1С есть в ERP, суммы/даты/НДС/договоры/заявки совпадают, сальдо на начало и конец совпадает.",
                "MariaDB ERP + существующий SOAP-сервис 1С.",
                "Протокол запуска сверки, список расхождений, итоговый статус по документам и сальдо.",
            ),
            (
                "Фоновый обработчик",
                "Считать тяжелые исторические сверки по расписанию или по ручному запуску без блокировки интерфейса.",
                "Те же источники, но в пакетном режиме с сохранением снапшота результата.",
                "Готовый результат сверки/выгрузки, который открывается из истории запусков.",
            ),
        ],
        [1700, 3300, 2620, 1740],
        font_size=7.25,
    )
    add_callout(
        doc,
        "Терминология",
        "`Сальдо по поставкам` - ERP-расчет по выбранным поставкам. `Монитор ERP ↔ 1С` - отдельный процесс сравнения двух систем. На странице сальдо нельзя показывать документные статусы 1С: там остаются только `Долг` и `Переплата`.",
        fill=WARN_FILL,
    )


def add_workflow(doc):
    add_heading(doc, "2. Рабочий процесс", 1)
    steps = [
        "Бухгалтер открывает страницу `Сальдо по поставкам` из ERP: из клиента, договора, поставки или общего раздела.",
        "Выбирает клиента, ЮЛ, договор, период или конкретный набор поставок. ERP считает только выбранный срез, без обращения к 1С.",
        "Экран показывает агрегированную иерархию и ERP-сальдо: счета, оплаты, возмещаемые расходы, невозмещаемые расходы, долг или переплату.",
        "Кнопка `Выгрузить Excel` формирует бухгалтерскую форму с объединенными ячейками по поставке.",
        "Если бухгалтеру нужна проверка учетных систем, он открывает отдельную страницу `Монитор ERP ↔ 1С`.",
        "На странице сверки выбирает ЮЛ, договор, период и при необходимости поставку; нажимает `Запустить сверку`.",
        "Backend создает запуск сверки, собирает ERP-срез, получает 1С-срез через SOAP, сравнивает документы и сальдо, сохраняет протокол.",
        "Пользователь видит итог: совпало ли сальдо, сколько документов совпало, чего нет в 1С, чего нет в ERP, какие реквизиты или суммы расходятся.",
        "Фоновые обработчики могут запускать тяжелую историческую сверку по расписанию и сохранять результат для просмотра без пересчета на экране.",
    ]
    for step in steps:
        add_number(doc, step)


def add_functional_requirements(doc):
    add_heading(doc, "3. Функциональные требования", 1)
    add_table(
        doc,
        ["ID", "Требование", "Критерий выполнения"],
        [
            (
                "FR-01",
                "Страница `Сальдо по поставкам` открывается из карточки клиента, договора, поставки или общего раздела. В URL допускаются параметры `contact_id`, `client_id`, `dog_id`, `spec_id`, `period`, `limit`, `scope`.",
                "При `client_id + scope=legal` показывается одно ЮЛ; при `contact_id/scope=contact` - все ЮЛ клиента.",
            ),
            (
                "FR-02",
                "Система строит иерархию `Клиент -> ЮЛ -> Договор -> Поставка -> Документы` только по исходным таблицам ERP.",
                "На родительских уровнях показаны агрегаты дочерних поставок, документы раскрываются только внутри поставки.",
            ),
            (
                "FR-03",
                "Система выбирает операции поставки по прямой связи `f_parenttype=2` и альтернативной связи `f_parenttype=4` через `veda_categs`.",
                "Операции с `f_parenttype=4` не теряются и участвуют в счетах, оплатах и расходах.",
            ),
            (
                "FR-04",
                "Система считает оплаты по распределенной сумме `veda_acchist_docs.f_clssum` и контролирует ее против `get_paidsum`.",
                "Полная сумма банковского документа не используется как сумма оплаты поставки.",
            ),
            (
                "FR-05",
                "Система считает возмещаемые и невозмещаемые суммы через действующую ERP-процедуру `get_realizsum(oper_id)` с разделением по `f_isvozm`.",
                "`f_isvozm=1` попадает в возмещаемые расходы, `f_isvozm=2` - в невозмещаемые, неизвестное значение получает ошибку данных.",
            ),
            (
                "FR-06",
                "Страница `Монитор ERP ↔ 1С` запускается отдельно от матрицы сальдо и работает через текущий PHP ERP и существующий сервис 1С.",
                "Результат страницы - протокол расхождений, а не бухгалтерская сальдовка по поставкам.",
            ),
            (
                "FR-07",
                "Сверка ERP ↔ 1С возвращает статусы `ОК`, `Нет в 1С`, `Нет в ERP`, `Сумма расходится`, `Дата расходится`, `Номер расходится`, `Договор/заявка расходится`, `НДС расходится`, `Нет СФ/УПД`, `Сальдо расходится`, `Ошибка источника`.",
                "Пользователь видит причину, ERP-значение, 1С-значение, конкретный вид расхождения и ссылку на исходный документ.",
            ),
            (
                "FR-08",
                "Кнопка `Скопировать` формирует TSV, где значения объединенных ячеек повторены в каждой строке.",
                "Вставка в Excel не теряет контекст поставки.",
            ),
            (
                "FR-09",
                "Кнопка `Выгрузить Excel` формирует бухгалтерский XLSX, не копию экрана.",
                "Колонки и объединения соответствуют разделу `Excel-выгрузка`; полный список документов выводится через Enter.",
            ),
            (
                "FR-10",
                "Каждый запуск страницы `Сверка ERP ↔ 1С` пишет журнал запуска и журнал расхождений.",
                "По журналу можно посчитать частоту проблем по клиенту, договору, поставке, типу операции и коду ошибки.",
            ),
            (
                "FR-11",
                "Фоновый обработчик может выполнять тяжелые сверки по расписанию или вручную.",
                "Пользователь видит историю запусков, статус выполнения, дату актуальности данных и ссылку на результат.",
            ),
            (
                "FR-12",
                "Сверка ERP ↔ 1С должна сравнивать не только документы, но и сальдо.",
                "Итоговый статус включает начальное сальдо, обороты за период, конечное сальдо и документные расхождения.",
            ),
        ],
        [850, 5450, 3060],
        font_size=7.3,
    )


def add_reconciliation_page_requirements(doc):
    add_heading(doc, "3.1 Отдельная страница `Монитор ERP ↔ 1С`", 2)
    add_body(
        doc,
        "Эта страница нужна не для подготовки клиентской Excel-сальдовки, а для контроля качества интеграции и учета: все ли документы ERP выгружены в 1С, нет ли лишних документов в 1С, совпадают ли суммы, реквизиты, НДС и сальдо.",
    )
    add_table(
        doc,
        ["Блок страницы", "Требование"],
        [
            (
                "Фильтры запуска",
                "ЮЛ, организация, договор, период `date_from/date_to`; опционально поставка или список поставок. Для полноценного акта сверки обязателен непрерывный период, а не произвольный набор поставок.",
            ),
            (
                "Кнопка запуска",
                "`Запустить сверку` создает `run_id`. Если объем малый - расчет синхронный; если период большой или много договоров - фоновая задача.",
            ),
            (
                "Сводка результата",
                "Показать `Сальдо ERP`, `Сальдо 1С`, `Расхождение`, количество документов ERP, количество документов 1С, сколько совпало, сколько отсутствует, сколько расходится.",
            ),
            (
                "Вкладка `Расхождения`",
                "Таблица строк, где есть проблема: тип документа, ERP-документ, 1С-документ, поле расхождения, значение ERP, значение 1С, сумма разницы, действие.",
            ),
            (
                "Вкладка `ERP без 1С`",
                "Документы, которые есть в ERP и должны быть в 1С, но не найдены: счет, оплата, акт, УПД/СФ, договор/поставка.",
            ),
            (
                "Вкладка `1С без ERP`",
                "Документы, которые вернула 1С, но ERP-аналог не найден. Это обязательный блок: без него сверка односторонняя.",
            ),
            (
                "Вкладка `Сальдо`",
                "Начальное сальдо ERP/1С, обороты ERP/1С, конечное сальдо ERP/1С. Если выбран произвольный набор поставок, вместо акта показывать предупреждение: `Нельзя рассчитать начальное сальдо для произвольного набора поставок`.",
            ),
            (
                "История запусков",
                "Список запусков с параметрами, временем, пользователем, статусом, длительностью, количеством расхождений и ссылкой на результат.",
            ),
        ],
        [2300, 7060],
        font_size=7.75,
    )
    add_table(
        doc,
        ["Что сравниваем", "ERP-сторона", "1С-сторона", "Статус при расхождении"],
        [
            ("Начальное сальдо", "ERP-оборотка до `date_from` по ЮЛ/договору", "1С-сальдо на начало периода", "`OPENING_BALANCE_MISMATCH`"),
            ("Обороты периода", "ERP-документы и движения за период", "1С-документы и движения за период", "`TURNOVER_MISMATCH`"),
            ("Конечное сальдо", "начальное ERP + обороты ERP", "начальное 1С + обороты 1С", "`CLOSING_BALANCE_MISMATCH`"),
            ("Счет покупателю", "`veda_schets` покупателя", "1С счет покупателю", "`MISSING_1C`, `MISSING_ERP`, `AMOUNT_MISMATCH`, `DATE_MISMATCH`"),
            ("Оплата", "`veda_acchist_docs.f_clssum` + банковский документ", "1С поступление/списание и расшифровка", "`PAYMENT_MISSING`, `PAYMENT_AMOUNT_MISMATCH`"),
            ("Акт/УПД/СФ", "`veda_akts` и детали", "1С реализация/поступление/ЭДО", "`NO_SF_UPD`, `VAT_MISMATCH`, `AMOUNT_MISMATCH`"),
            ("Договор/поставка", "`veda_dogs`, `veda_specs`, коды 1С", "1С договоры контрагентов", "`CONTRACT_MISMATCH`, `SPEC_KEY_MISSING`"),
        ],
        [1500, 2520, 2300, 3040],
        font_size=7.1,
    )
    add_callout(
        doc,
        "Ключевое отличие от матрицы сальдо",
        "Матрица сальдо отвечает на вопрос `сколько клиент должен/переплатил по выбранным поставкам`. Монитор ERP ↔ 1С отвечает на вопрос `совпадают ли две учетные системы`. Это разные интерфейсы, разные статусы и разные критерии приемки.",
        fill=LIGHT_FILL,
    )


def add_runner_requirements(doc):
    add_heading(doc, "3.2 Фоновый обработчик для тяжелых сверок", 2)
    add_body(
        doc,
        "Историческая сверка по всем клиентам и всем договорам не должна выполняться как обычный HTTP-запрос экрана. Для этого нужен фоновый обработчик, который запускается по расписанию или вручную и сохраняет результат.",
    )
    add_table(
        doc,
        ["Требование", "Описание"],
        [
            ("Режимы запуска", "ручной запуск по клиенту/договору/периоду; ночной запуск по всем активным клиентам; повторный запуск только проблемных договоров"),
            ("Сохранение снапшота", "сохранять параметры запуска, ERP-срез, 1С-срез, итоговые суммы и строки расхождений; экран открывает сохраненный результат"),
            ("Ограничение нагрузки", "batch по клиентам/договорам, лимит параллельности, таймаут SOAP, повтор при временной ошибке 1С"),
            ("Актуальность", "у каждого результата показывать дату расчета и источник: ERP, 1С, сохраненный результат, ошибка источника"),
            ("Повторяемость", "один и тот же `run_id` должен хранить использованные входные данные, чтобы итог можно было объяснить позже"),
            ("Экспорт", "XLSX по результату фонового обработчика формируется из сохраненного результата, а не пересчитывает тяжелые SQL и SOAP заново"),
        ],
        [2300, 7060],
        font_size=7.85,
    )


def add_nonfunctional_requirements(doc):
    add_heading(doc, "4. Нефункциональные требования", 1)
    add_table(
        doc,
        ["Группа", "Требование"],
        [
            (
                "Производительность",
                "На экране нельзя запускать расчет по всему клиенту без лимита. Онлайн-экран должен иметь малый лимит по умолчанию и пагинацию: контрольная проверка показала, что 3 поставки считаются десятки секунд, а 10 поставок уже неприемлемы как обычный HTTP-запрос. Тяжелые процедуры `get_paidsum/get_realizsum/get_expensessum/get_profit` выполнять только по выбранным `spec_id`. Большие выборки - только через фоновый обработчик/кэш результата.",
            ),
            (
                "Надежность",
                "Недоступность 1С должна фиксироваться как ошибка источника, а не как массовый статус `Нет в 1С`. Повторный запуск с теми же параметрами должен давать тот же расчет при неизменных данных.",
            ),
            (
                "Безопасность",
                "Доступ к странице и XLSX определяется правами текущей ERP. Логины, пароли, SOAP URL и параметры подключения к MariaDB не передаются во frontend и не попадают в XLSX.",
            ),
            (
                "Совместимость",
                "Production-реализация остается в текущем PHP ERP + MariaDB + сервис 1С. Внешние аналитические файлы не являются источником расчета.",
            ),
            (
                "Наблюдаемость",
                "Каждый запуск сохраняет `run_id`, пользователя, параметры фильтра, время выполнения SQL и SOAP, количество поставок, количество документов, количество расхождений и техническую ошибку при сбое.",
            ),
            (
                "Поддерживаемость",
                "Коды статусов, типы документов, значения `f_isvozm`, правила `f_type` счетов и соответствия SOAP-полей должны быть вынесены в именованные константы/справочники, а не хардкодиться текстом в шаблоне.",
            ),
            (
                "Точность",
                "Денежное сравнение выполняется с допуском 0,01. Суммы разных валют не агрегируются без утвержденного правила конвертации; интерфейс показывает причину `Смешанные валюты`.",
            ),
            (
                "Аудит",
                "XLSX и экран должны строиться из одного нормализованного набора данных. Нельзя, чтобы экран и Excel считали одни и те же поля разными SQL-путями.",
            ),
        ],
        [1800, 7560],
        font_size=7.65,
    )


def add_hierarchy(doc):
    add_heading(doc, "5. Иерархия данных", 1)
    add_body(
        doc,
        "Матрица строится как дерево. Верхние уровни являются агрегатами дочерних поставок. На них нельзя механически повторять сумму выбранной строки документа: это и дает ошибочную картину, когда у клиента оплата больше счета, а дельта выглядит нулевой.",
    )
    add_table(
        doc,
        ["Уровень", "ERP-таблица", "Связь", "Что выводить"],
        [
            ("Клиент", "veda_contacts", "veda_contacts.f_id = veda_clients.f_contactid", "итоги по всем ЮЛ клиента"),
            ("ЮЛ", "veda_clients", "veda_clients.f_id = veda_dogs.f_contrid", "итоги по юридическому лицу"),
            ("Договор", "veda_dogs", "veda_dogs.f_id = veda_specs.f_dogid", "итоги по договору"),
            ("Поставка / спецификация", "veda_specs", "veda_specs.f_id", "счета, оплаты, расходы, СФ/УПД, дельта"),
            ("Документы внутри поставки", "veda_schets, veda_acchist_docs, veda_akts", "операция поставки / `f_operid`", "раскрытие только для диагностики"),
        ],
        [1150, 1750, 3250, 3210],
    )
    add_code(
        doc,
        """
-- Клиент -> ЮЛ -> договор -> поставка
SELECT c.f_id contact_id, cl.f_id client_id, d.f_id dog_id, s.f_id spec_id
FROM veda_contacts c
JOIN veda_clients cl ON cl.f_contactid = c.f_id
JOIN veda_dogs d ON d.f_contrid = cl.f_id
JOIN veda_specs s ON s.f_dogid = d.f_id
WHERE c.f_id = :contact_id;
""",
    )


def add_ui(doc):
    add_heading(doc, "6. Требования к интерфейсу сальдо по поставкам", 1)
    add_body(
        doc,
        "Этот экран остается полезным бухгалтерским инструментом, но не является финальным протоколом сверки ERP ↔ 1С. Он показывает ERP-сальдо по выбранным поставкам и готовит XLSX для работы с клиентом.",
    )
    add_table(
        doc,
        ["Элемент", "Требование"],
        [
            ("Первая колонка", "`Иерархия` закреплена слева при горизонтальном скролле."),
            ("Деньги", "Все денежные колонки выровнены вправо; формат с валютой сохраняется."),
            ("Родительские строки", "Клиент, ЮЛ, договор и поставка выделены жирнее и показывают агрегаты."),
            ("Документы", "Счета, оплаты, акты и СФ/УПД раскрываются внутри поставки, а не показываются сразу всем списком."),
            ("Статусы", "На странице сальдо показывать только `Долг` и `Переплата`. Документные статусы 1С (`Нет в 1С`, `Сумма расходится`, `Договор/заявка расходится`) показывать только на странице `Монитор ERP ↔ 1С`."),
            ("Поиск", "По клиенту, ЮЛ, ИНН, `contact_id`, `client_id`, договору, заявке/спецификации, счету, акту, СФ/УПД. `client_id + scope=legal` открывает одно ЮЛ, `contact_id/scope=contact` - все ЮЛ клиента."),
        ],
        [1900, 7460],
    )
    add_image(doc, IMG_SALDO, "Рисунок 1. Референсный вид экрана `Сальдо по поставкам`: несколько поставок договора 660/1, ERP-расчет без обращения к 1С.")
    add_image(doc, IMG_MONITOR, "Рисунок 2. Референсный вид страницы `Монитор ERP ↔ 1С`: проверка документов по нескольким поставкам с источником 1С.")
    add_image(doc, IMG_MONITOR_ISSUES, "Рисунок 3. Режим `Только расхождения`: пример статуса, когда документ найден, но договор/заявка не совпадает.")


def add_mapping(doc):
    add_heading(doc, "7. Маппинг ERP и 1С", 1)
    add_body(
        doc,
        "Доработка должна использовать текущий PHP-код ERP и существующий 1С-сервис. Ниже перечислены сущности, которые нужно сравнивать в разрезе поставки.",
    )
    add_table(
        doc,
        ["ERP", "1С", "Ключ / поля сверки", "Колонка или статус"],
        [
            ("veda_contacts / veda_clients", "Справочник.Контрагенты", "код 1С, ИНН, КПП, название, признак удаления", "иерархия, справочник"),
            ("veda_dogs", "Справочник.ДоговорыКонтрагентов", "код 1С, номер, дата, организация, контрагент, валюта", "договор"),
            ("veda_specs + veda_spr", "договор-заявка / спецификация", "veda_specs.f_id, f_num, f_typez, f_subtype, f_kod1cb, f_kod1cp; тип через veda_spr f_type=33, подтип через f_type=130", "№ спецификации / заявки; view_specinv/view_specs не использовать"),
            ("veda_schets, f_type=1", "Документ.СчетНаОплатуПокупателю", "номер/код 1С, дата, организация, контрагент, договор, сумма", "Счет, Сумма по счету"),
            ("veda_acchist_docs", "Поступление/списание по банку", "f_docid, f_doctype=3, f_clssum, код/дата 1С", "Сумма оплаты"),
            ("veda_akts + veda_akts_details_opers", "Поступление/реализация, строки услуг", "акт, строка, операция, сумма, ставка НДС, договор строки", "Возм./невозм. расходы, НДС"),
            ("veda_akts / СФ-УПД", "Реализация, УПД, СФ", "номер/дата, сумма, НДС, контрагент, договор", "№ счф, Нет СФ/УПД"),
            ("1С акт сверки", "Документ.АктСверкиВзаиморасчетов", "организация, контрагент, договор, сальдо, обороты", "контроль итогов"),
        ],
        [1700, 1900, 3860, 1900],
        font_size=7.8,
    )
    add_heading(doc, "7.1 SOAP-методы 1С, которые нужны", 2)
    add_table(
        doc,
        ["Метод 1С", "Использование в сверке"],
        [
            ("getClients", "справочник контрагентов, договоры `dogs`, счета `accs`, признаки удаления"),
            ("getOrgs", "справочник организаций"),
            ("getAcchist", "банковские документы и расшифровка платежа по счетам/договорам"),
            ("getAkts", "акты реализации и поступления за период"),
            ("getEDOArh / getEDPacket", "СФ/УПД и ЭДО-статусы, если закрывающий документ в ЭДО"),
            ("getcoacsu / getcoacsuinfo", "итоги 1С акта сверки по контрагенту без договора и детализация по договорам"),
        ],
        [2050, 7310],
        font_size=8.2,
    )
    add_heading(doc, "7.2 Что передаем и что сравниваем", 2)
    add_table(
        doc,
        ["Блок", "ERP передает / выбирает", "1С возвращает", "Сравнение"],
        [
            ("Контрагент", "`contact_id`, `client_id`, ИНН, КПП, код 1С ЮЛ", "`getClients`: код, ИНН, КПП, название, удален", "контрагент найден и не удален; ИНН/КПП не конфликтуют"),
            ("Договор", "`dog_id`, номер, дата, организация, `dog_code1c`", "`getClients.dogs` или `getcoacsuinfo`: код договора, номер, организация", "код/номер договора и организация совпадают"),
            ("Поставка", "`veda_specs.f_id`, `f_num`, `f_typez`, `f_subtype`, `f_kod1cb/f_kod1cp`, расшифровка через `veda_spr`", "договоры-заявки/спецификации из 1С по кодам", "тип и номер выводятся из исходных таблиц ERP; полный ОК возможен только если код 1С найден"),
            ("Счет", "покупательские `veda_schets.f_type=1` по операциям поставки", "`getAcchist`/счет 1С в расшифровке или SOAP-документ счета", "номер, дата, сумма, валюта, организация, контрагент"),
            ("Оплата", "`veda_acchist_docs.f_clssum`, `f_docid`, `f_acchistid`", "`getAcchist`: банковский документ и расшифровка платежа", "сравнивается распределенная доля, а не полная сумма банковской выписки"),
            ("Реализация", "`get_realizsum(oper_id)` по операциям `f_isvozm=1/2`", "`getAkts`, `getEDOArh/getEDPacket`: акты, УПД/СФ, строки услуг", "сумма, НДС, закрывающий документ, договор строки"),
        ],
        [1200, 2830, 2530, 2800],
        font_size=7.45,
    )


def add_field_dictionary(doc):
    add_heading(doc, "8. Словарь полей и маппинг колонок", 1)
    add_body(
        doc,
        "Этот раздел нужен разработчику как прямое соответствие между пользовательской колонкой, ERP-источником, 1С-источником, формулой и статусом. Если поле не найдено в 1С, оно не должно молча заменяться похожим полем.",
    )
    add_table(
        doc,
        ["Поле UI/XLSX", "ERP-источник", "1С-источник", "Правило заполнения"],
        [
            (
                "Иерархия",
                "`veda_contacts`, `veda_clients`, `veda_dogs`, `veda_specs`",
                "`getClients`, договоры `dogs`, `getcoacsuinfo`",
                "Клиент и ЮЛ группируют поставки; договор и поставка берутся из ERP, 1С подтверждает код/договор.",
            ),
            (
                "№ спецификации",
                "`veda_specs.f_num`, `f_typez`, `f_subtype`, `f_dt`, `f_tovar`; справочники `veda_spr(f_type=33/130)`",
                "договор-заявка/спецификация из 1С по коду, если код есть",
                "Подпись формируется как `{тип} №{номер}`: `Заявка №1063`, `Спецификация №...`; `Спецификация` не хардкодить.",
            ),
            (
                "Счет",
                "`veda_schets.f_num`, `f_kod1c`, `f_dt`, `f_sum`, `f_val`, связь с операцией",
                "1С счет покупателю: номер, дата, сумма, организация, контрагент",
                "Показывать только счета покупателю. Если есть агрегирующий счет, основной счет выводится один раз.",
            ),
            (
                "Сумма по счету",
                "`veda_schets.f_sum` по покупательскому счету",
                "сумма счета 1С",
                "Сверка счета отдельная; в формулу долга/переплаты не входит.",
            ),
            (
                "Сумма оплаты",
                "`SUM(veda_acchist_docs.f_clssum)` и контроль `SUM(get_paidsum)`",
                "`getAcchist`: банковский документ и расшифровка платежа",
                "Считать только распределенную долю оплаты по операции/счету, не полную сумму банковского документа.",
            ),
            (
                "Возмещаемые расходы",
                "`SUM(get_realizsum(oper_id)) WHERE veda_spec_invoices.f_isvozm = 1`",
                "`getAkts`/ЭДО: акты/УПД/строки 1С",
                "Сумма клиентской реализации по возмещаемым операциям.",
            ),
            (
                "Невозмещаемые расходы",
                "`SUM(get_realizsum(oper_id)) WHERE veda_spec_invoices.f_isvozm = 2`",
                "`getAkts`/ЭДО: акты/УПД/строки 1С",
                "Сумма клиентской реализации по невозмещаемым операциям. Подрядные расходы без клиентской реализации исключаются из сальдо.",
            ),
            (
                "№ счф",
                "`veda_akts`, основной акт `f_mainakt`, код/номер/дата/НДС закрывающего документа",
                "`getAkts`, `getEDOArh`, `getEDPacket`",
                "Вывести все закрывающие документы через Enter. Не заменять список текстом `+N документов`.",
            ),
            (
                "(+/-)",
                "`SUM(get_paidsum) - SUM(get_realizsum)`",
                "`getcoacsuinfo` может использоваться как контроль итогов",
                "Отрицательное значение показывать как `Долг`, положительное как `Переплата`; термин `Остаток` не использовать.",
            ),
        ],
        [1450, 3030, 2350, 2530],
        font_size=7.0,
    )
    add_heading(doc, "8.1 Маппинг журнала расхождений", 2)
    add_table(
        doc,
        ["Поле журнала", "Источник", "Зачем нужно"],
        [
            ("run_id", "создается backend при запуске сверки", "связать строки одного запуска"),
            ("contact_id, client_id, dog_id, spec_id", "ERP-контекст", "группировка проблем по клиенту, ЮЛ, договору и поставке"),
            ("operation_id, operation_type_id, operation_type_name", "`veda_spec_invoices`, `veda_typeopers`", "понять, какие операции чаще дают расхождения"),
            ("doc_type, erp_doc_id, erp_doc_key", "ERP-документ: счет, платеж, акт, УПД/СФ", "идентифицировать проблемный документ"),
            ("onec_doc_type, onec_doc_key, onec_guid", "SOAP-ответ 1С", "идентифицировать документ 1С"),
            ("mismatch_code", "результат сравнения", "агрегировать проблемы по типам"),
            ("erp_value, onec_value, diff_amount", "расчет сравнения", "видеть точную разницу"),
            ("f_isvozm, nds_erp, nds_1c", "операция ERP и документ 1С", "отделить возмещаемость от НДС-расхождения"),
            ("source_table, source_row_id", "таблица и ID строки ERP", "быстрый переход к исходной записи"),
        ],
        [2450, 3600, 3310],
        font_size=7.55,
    )


def add_sources_and_selects(doc):
    add_heading(doc, "9. ERP-селекты и правила расчета", 1)
    add_heading(doc, "9.0 Ограничения выборки", 2)
    add_table(
        doc,
        ["Ограничение", "Требование"],
        [
            ("Фильтр обязателен", "Запуск без клиента/периода/договора запрещен. Минимум: клиент + период; для первичного экрана желательно договор."),
            ("Размер страницы", "Онлайн-экран показывает небольшой срез поставок: рекомендуемый стартовый лимит 3-5, все остальное через пагинацию или фоновую XLSX-задачу. Проверка 10 поставок на договоре 660/1 не уложилась в приемлемое время."),
            ("Поиск по ЮЛ", "Фильтр должен искать по `veda_contacts`, `veda_clients`, ИНН, `contact_id`, `client_id`, договору, поставке, счету и номеру 1С."),
            ("Тяжелые процедуры", "`get_paidsum`, `get_realizsum`, `get_expensessum`, `get_profit` считать только по выбранным `spec_id` и их операциям, не по всему договору без лимита."),
            ("SOAP 1С", "Вызовы 1С выполнять батчами по организациям/периодам; не дергать 1С отдельно на каждую строку таблицы, если можно получить период одним вызовом."),
            ("Документы внутри поставки", "На экране показывать агрегаты; полный список актов/СФ/УПД выводить в раскрытии или XLSX, чтобы не раздувать матрицу."),
        ],
        [2100, 7260],
        font_size=8,
    )
    add_heading(doc, "9.1 Список поставок по клиенту/ЮЛ/договору", 2)
    add_code(
        doc,
        """
SELECT
    s.f_id AS spec_id,
    s.f_num AS spec_num,
    COALESCE(NULLIF(spec_type.f_dopprstr, ''), NULLIF(spec_type.f_name, ''), '') AS spec_type,
    COALESCE(NULLIF(spec_subtype.f_dopprstr, ''), NULLIF(spec_subtype.f_name, ''), '') AS spec_subtype,
    DATE_FORMAT(s.f_dt, '%Y-%m-%d') AS spec_date,
    s.f_tovar AS spec_name,
    d.f_id AS dog_id,
    d.f_dogname AS dog_number,
    d.f_kod1c AS dog_code1c,
    cl.f_id AS legal_id,
    cl.f_cname AS legal_name,
    cl.f_inn AS legal_inn,
    contact.f_id AS contact_id,
    COALESCE(contact.f_cname, contact.f_name, cl.f_cname) AS contact_name
FROM veda_specs s
JOIN veda_dogs d ON d.f_id = s.f_dogid
JOIN veda_clients cl ON cl.f_id = d.f_contrid
LEFT JOIN veda_contacts contact ON contact.f_id = cl.f_contactid
LEFT JOIN veda_spr spec_type ON spec_type.f_type = 33 AND spec_type.f_num = s.f_typez
LEFT JOIN veda_spr spec_subtype ON spec_subtype.f_type = 130 AND spec_subtype.f_num = s.f_subtype
WHERE
  (
       (:scope = 'legal' AND cl.f_id = :client_id)
    OR (:scope = 'contact' AND contact.f_id = :contact_id)
    OR (:scope = 'mixed' AND (cl.f_id = :client_id OR contact.f_id = :client_id))
  )
  AND (:dog_id IS NULL OR d.f_id = :dog_id)
  AND s.f_dt BETWEEN :date_from AND :date_to
ORDER BY s.f_dt DESC, s.f_id DESC
LIMIT :limit OFFSET :offset;
""",
    )
    add_heading(doc, "9.2 Операции поставки", 2)
    add_code(
        doc,
        """
-- Операции, привязанные к поставке напрямую или через f_parenttype=4
SELECT si.*
FROM veda_spec_invoices si
LEFT JOIN veda_categs cg
  ON cg.f_objectid = si.f_id
 AND cg.f_ctgtype = 24
 AND cg.f_objecttype = 5
WHERE (si.f_parenttype = 2 AND si.f_specid = :spec_id)
   OR (si.f_parenttype = 4 AND CAST(cg.f_valstr AS SIGNED) = :spec_id);
""",
    )
    add_heading(doc, "9.3 Счета покупателю", 2)
    add_code(
        doc,
        """
-- В матрицу идут только счета покупателю
SELECT main.f_id, main.f_num, main.f_sum, main.f_val, main.f_kod1c, main.f_dt
FROM veda_schets child
JOIN veda_schets main
  ON main.f_id = IF(child.f_maininv > 0, child.f_maininv, child.f_id)
WHERE child.f_operid IN (:operation_ids)
  AND main.f_type = 1; -- 1 = счет покупателю, 2 = счет от поставщика
""",
    )
    add_heading(doc, "9.4 Оплаты", 2)
    add_code(
        doc,
        """
-- Нельзя брать всю сумму банковского документа: нужна только распределенная доля
SELECT ahd.f_docid operation_id, SUM(ahd.f_clssum) paid_sum
FROM veda_acchist_docs ahd
JOIN veda_acchist ah ON ah.f_id = ahd.f_acchistid
WHERE ahd.f_doctype = 3
  AND ahd.f_docid IN (:operation_ids)
  AND ah.f_type = 0
GROUP BY ahd.f_docid;
""",
    )
    add_heading(doc, "9.5 Расходы", 2)
    add_table(
        doc,
        ["Тип", "ERP-признак", "Как считать", "Комментарий"],
        [
            ("Возмещаемые", "f_isvozm = 1", "`get_realizsum` по реализации/актам и деталям операции", "попадают в колонку `Возмещаемые расходы`"),
            ("Невозмещаемые", "f_isvozm = 2", "`get_realizsum` по реализации/актам и деталям операции", "попадают в колонку `Невозмещаемые расходы`"),
            ("Расходы подрядчика", "подрядная себестоимость без клиентской реализации", "не включать в взаиморасчет с клиентом", "логировать как `EXCLUDED_CONTRACTOR_EXPENSE`"),
        ],
        [1750, 2100, 2950, 2560],
        font_size=8,
    )
    add_callout(
        doc,
        "Ограничение по агентской деятельности",
        "Если услуга подрядчика отражена в 1С как расходы компании на вкладке `Услуги` с субконто `Агентское вознаграждение / Услуги сторонних организаций`, договор клиента нельзя искусственно подтягивать в взаиморасчет. Такая строка подтверждает себестоимость, но не должна менять сальдо клиента.",
        fill=WARN_FILL,
    )
    add_heading(doc, "9.6 Закрывающие документы", 2)
    add_code(
        doc,
        """
-- Прямые акты по операции
SELECT
    akt.f_id,
    akt.f_kod1c,
    akt.f_num,
    akt.f_dt,
    akt.f_dt1c,
    akt.f_sum,
    akt.f_val,
    akt.f_type,
    akt.f_mainakt,
    oper.f_id AS operation_id,
    oper.f_isvozm,
    oper.f_nds
FROM veda_spec_invoices oper
JOIN veda_akts akt ON akt.f_operid = oper.f_id
WHERE oper.f_id IN (:operation_ids)

UNION

-- Акты через распределенные строки детализации
SELECT
    akt.f_id,
    akt.f_kod1c,
    akt.f_num,
    akt.f_dt,
    akt.f_dt1c,
    COALESCE(ado.f_sum, det.f_sum + det.f_ndssum) AS f_sum,
    akt.f_val,
    akt.f_type,
    akt.f_mainakt,
    oper.f_id AS operation_id,
    oper.f_isvozm,
    oper.f_nds
FROM veda_spec_invoices oper
JOIN veda_akts_details_opers ado ON ado.f_operid = oper.f_id
JOIN veda_akts_details det ON det.f_id = ado.f_akts_detailsid
JOIN veda_akts akt ON akt.f_id = det.f_aktid
WHERE oper.f_id IN (:operation_ids);
""",
    )
    add_heading(doc, "9.7 Сквозной алгоритм для PHP-разработчика", 2)
    add_body(
        doc,
        "Матрица должна строиться не из Excel и не из внешнего файла. Источник экрана и XLSX - один нормализованный набор блоков поставок. Ниже указано, какие данные собрать, что агрегировать и в какие колонки положить.",
    )
    add_table(
        doc,
        ["Шаг", "Данные / запрос", "Что сохранить в блоке поставки"],
        [
            (
                "1. Контекст",
                "`veda_contacts -> veda_clients -> veda_dogs -> veda_specs`; тип поставки через `veda_spr(f_type=33)`, подтип через `veda_spr(f_type=130)`; фильтры `contact_id`, `client_id`, `dog_id`, период или список `spec_id/f_num`.",
                "Клиент, ЮЛ, договор, поставка: id, название, ИНН, номер договора, код 1С, `spec_type`, `spec_subtype`, номер/дата поставки.",
            ),
            (
                "2. Операции",
                "`veda_spec_invoices` по `f_parenttype=2` и по `f_parenttype=4` через `veda_categs(f_ctgtype=24, f_objecttype=5)`.",
                "`operations[]`: `oper_id`, тип операции, `f_isvozm`, НДС, `get_paidsum`, `get_realizsum`, `get_expensessum`, `get_profit`.",
            ),
            (
                "3. Счета",
                "Из `fetch_erp_docs(spec_id)` брать `doc_kind='schet'`, `type_name='Счет покупателю'`, `invoice_id=0`; при отдельном SQL - только `veda_schets.f_type=1` и основной счет.",
                "`customer_invoices[]`: номер, дата, сумма, валюта, код 1С. Счета поставщиков не добавлять.",
            ),
            (
                "4. Оплаты",
                "`SUM(veda_acchist_docs.f_clssum)` по `f_doctype=3`, `f_docid IN operation_ids`, `JOIN veda_acchist ON f_acchistid`, `ah.f_type=0`; контроль с `get_paidsum(oper_id)`.",
                "`paid_total_get_paidsum`, `paid_total_acchist_docs`; если различаются больше 0,01 - статус `PAYMENT_ROUTINE_VS_DOCS_MISMATCH`.",
            ),
            (
                "5. Реализация",
                "Для каждой операции взять `get_realizsum(oper_id)` и разделить по `veda_spec_invoices.f_isvozm`.",
                "`reimbursable = SUM(get_realizsum WHERE f_isvozm=1)`, `non_reimbursable = SUM(get_realizsum WHERE f_isvozm=2)`.",
            ),
            (
                "6. Закрывающие документы",
                "Из `fetch_erp_docs(spec_id)` брать `doc_kind='act'`; дедупликация по `code1c/number/date/sum`.",
                "`closing_docs[]`: каждый акт/УПД/СФ отдельной строкой; не сворачивать список в счетчик количества документов.",
            ),
            (
                "7. Дельта",
                "`delta = SUM(get_paidsum) - SUM(get_realizsum)`.",
                "`settlement_delta_paid_minus_realization`; ненулевое значение - статус `SETTLEMENT_BALANCE_NONZERO`, а не ошибка SQL.",
            ),
        ],
        [700, 4380, 4280],
        font_size=7.15,
    )
    add_table(
        doc,
        ["Нормализованный блок", "Обязательные поля"],
        [
            ("Контекст", "`contact_id/name/inn`, `client_id/name/inn`, `dog_id/number/code1c`, `spec_id/number/type/subtype/date/name/code1c`."),
            ("Счета покупателю", "список `number`, `date`, `sum`, `currency`, `code1c`; только `f_type=1`, основной счет без дублей дочерних строк."),
            ("Оплаты", "`paid_total_get_paidsum`, `paid_total_acchist_docs`, список распределенных платежей по `f_clssum`."),
            ("Реализация", "`reimbursable_realization_get_realizsum`, `non_reimbursable_realization_get_realizsum`, `realization_total_get_realizsum`."),
            ("Закрывающие документы", "все акты/УПД/СФ: `number`, `date`, `sum`, `code1c`, `nds_rate`; без сокращения до счетчика документов в XLSX."),
            ("Статусы", "для сальдо: только `Долг`/`Переплата`; для монитора ERP ↔ 1С: массив кодов расхождений с суммами ERP/1С, операцией, документом и причиной."),
        ],
        [2050, 7310],
        font_size=7.7,
    )
    add_table(
        doc,
        ["Колонка XLSX", "Заполнение", "Формат/объединение"],
        [
            ("A `№ спецификации`", "`{veda_spr.f_name/f_dopprstr по f_type=33} №{veda_specs.f_num}`; если тип пустой, `Поставка №{veda_specs.f_num}`.", "Объединить по числу строк счетов поставки."),
            ("C `Счет`", "Каждый `customer_invoices[].number` отдельной строкой.", "Не объединять, строки идут вертикально."),
            ("D `Сумма по счету`", "`customer_invoices[].sum + currency`.", "Не суммировать разные валюты без курса."),
            ("E `Сумма оплаты`", "`paid_total_get_paidsum`.", "Объединить по поставке; контроль с `paid_total_acchist_docs`."),
            ("G `Возмещаемые расходы`", "`reimbursable_realization_get_realizsum`.", "Объединить по поставке, деньги вправо."),
            ("H `Невозмещаемые расходы`", "`non_reimbursable_realization_get_realizsum`.", "Объединить по поставке, деньги вправо."),
            ("J `№ счф`", "`closing_docs` строками: `{code1c|number} от {date} ({sum})`.", "Объединить по поставке, перенос строки внутри ячейки."),
            ("K `(+/-)`", "`settlement_delta_paid_minus_realization`.", "Объединить по поставке; красный фон/текст при ненуле."),
            ("ИТОГО", "Суммы E/G/H/K по всем поставкам выборки.", "Подпись `ИТОГО`, без номеров строк внешних таблиц."),
        ],
        [1650, 5100, 2610],
        font_size=7.25,
    )
    add_body(
        doc,
        "Источник экрана и XLSX должен быть один: нормализованный набор блоков поставок, собранный из исходных таблиц ERP по правилам выше.",
    )


def add_column_rules(doc):
    add_heading(doc, "10. Колонки матрицы", 1)
    add_table(
        doc,
        ["Колонка", "ERP-источник", "Проверка с 1С / правило"],
        [
            ("Иерархия", "veda_contacts -> veda_clients -> veda_dogs -> veda_specs", "1С-справочники проверяют коды, ИНН, КПП и договоры."),
            ("№ спецификации", "veda_specs.f_typez/f_subtype/f_num/f_id/f_kod1cb/f_kod1cp + veda_spr(f_type=33/130)", "тип строки брать из исходных таблиц ERP: `Заявка`, `Спецификация`, `Сертификация` и т.д.; не хардкодить `Спецификация {num}`; не использовать view_specinv/view_specs."),
            ("Счет", "veda_schets.f_num, f_kod1c, f_dt, f_operid", "1С `СчетНаОплатуПокупателю`; `f_type=1` только покупательские счета."),
            ("Сумма по счету", "сумма счетов покупателю по поставке, с учетом валюты", "сверяется со счетом 1С, но не входит в формулу `(+/-)`."),
            ("Сумма оплаты", "SUM(veda_acchist_docs.f_clssum)", "сверяется с расшифровкой банковского документа 1С, не с общей суммой выписки."),
            ("Возмещаемые расходы", "`get_realizsum`, f_isvozm=1", "сверка по акту/УПД/поступлению и строке."),
            ("Невозмещаемые расходы", "`get_realizsum`, f_isvozm=2", "подрядные расходы без клиентской реализации не включать."),
            ("№ счф", "закрывающие документы ERP/1С", "все номера актов/СФ/УПД показывать через перенос строки."),
            ("(+/-)", "поступления - реализация", "сальдо поставки: дельта между распределенными поступлениями и реализацией (`get_realizsum`) по поставке."),
            ("Статус", "расчет сальдо ERP", "`Долг`, `Переплата` или пусто при нулевой дельте. Документные статусы 1С здесь не показывать."),
        ],
        [1280, 2940, 5140],
        font_size=7.7,
    )


def add_statuses(doc):
    add_heading(doc, "11. Статусы и сценарии проверки", 1)
    add_table(
        doc,
        ["Страница", "Статус", "Когда ставим", "Что сравниваем"],
        [
            ("Сальдо", "Долг", "`SUM(get_paidsum) - SUM(get_realizsum) < -0,01`", "только ERP-расчет по выбранным поставкам"),
            ("Сальдо", "Переплата", "`SUM(get_paidsum) - SUM(get_realizsum) > 0,01`", "только ERP-расчет по выбранным поставкам"),
            ("Сальдо", "пусто", "дельта в пределах допуска 0,01", "не показывать `ОК` и не показывать статусы 1С"),
            ("Монитор ERP ↔ 1С", "ОК", "ключ документа найден в ERP и 1С, суммы и обязательные реквизиты совпали", "код/номер/дата, организация, контрагент, договор, сумма, валюта, НДС при наличии"),
            ("Монитор ERP ↔ 1С", "Нет в 1С", "ERP-документ должен быть выгружен/найден, но 1С-сервис его не вернул", "тип документа, код 1С, номер, дата, организация"),
            ("Монитор ERP ↔ 1С", "Нет в ERP", "1С вернула документ, которому нет ERP-аналога", "номер/дата/организация/контрагент/договор/сумма"),
            ("Монитор ERP ↔ 1С", "Сумма расходится", "документ или агрегат найден, но сумма отличается больше допуска", "счет, оплата, акт, расход или итог поставки"),
            ("Монитор ERP ↔ 1С", "Дата расходится", "документ найден, но дата ERP и дата 1С различаются", "дата документа/движения"),
            ("Монитор ERP ↔ 1С", "Номер расходится", "документ найден по связке, но номер/код не совпадает", "номер ERP, код 1С, номер 1С"),
            ("Монитор ERP ↔ 1С", "Договор/заявка расходится", "сумма или номер найдены, но договор/заявка не совпадает", "dogcode, spec code, владелец, организация"),
            ("Монитор ERP ↔ 1С", "Нет СФ/УПД", "после успешной загрузки источников нет ожидаемого закрывающего документа", "наличие акта/УПД/СФ в ERP и 1С; не путать с возмещаемостью"),
            ("Монитор ERP ↔ 1С", "НДС расходится", "расход или акт найден, но ставка НДС ERP и ставка НДС 1С различаются", "только ставка НДС и сумма НДС; `f_isvozm` не является причиной этого статуса"),
            ("Монитор ERP ↔ 1С", "Нет ключа 1С", "ERP-документ есть, но нет кода/даты 1С для надежного матчинга", "f_kod1c, f_dt1c, номер, дата"),
            ("Монитор ERP ↔ 1С", "Сальдо расходится", "начальное, обороты или конечное сальдо ERP и 1С отличаются", "только для отдельной страницы сверки ERP ↔ 1С"),
            ("Оба", "Смешанные валюты", "в счетах есть разные валюты без утвержденной конвертации", "показывать суммы отдельно по валютам"),
            ("Оба", "Корректировки", "у поставки есть корректировки, которые меняют срез отчета", "основной отчет и корректировки показывать отдельными уровнями"),
        ],
        [1450, 1600, 3270, 3040],
        font_size=6.85,
    )
    add_callout(
        doc,
        "Запрет на смешение статусов",
        "На странице `Сальдо по поставкам` не показывать количество совпавших документов 1С и не использовать статус `ОК`. Документные проверки и причины расхождений показываются только на странице `Монитор ERP ↔ 1С`.",
        fill=WARN_FILL,
    )


def add_excel(doc):
    add_heading(doc, "12. Excel-выгрузка сальдо по поставкам", 1)
    add_body(
        doc,
        "Excel этого экрана не является протоколом сверки ERP ↔ 1С. Это бухгалтерская форма сальдо по поставкам: счета и оплаты идут строками, а расходы, СФ/УПД и дельта объединяются по количеству строк счетов внутри поставки.",
    )
    add_table(
        doc,
        ["Колонки Excel", "Правило"],
        [
            ("№ спецификации", "объединить по строкам счетов; текст вида `{тип ERP} №{номер}`, например `Заявка №1063`; тип брать из `veda_specs.f_typez/f_subtype` через `veda_spr`, без view"),
            ("Счет", "каждый покупательский счет отдельной строкой; если несколько счетов - переносы/строки"),
            ("Сумма по счету", "сумма счета в его валюте; не суммировать разные валюты без курса"),
            ("Сумма оплаты", "распределенные оплаты по счету/поставке; дополнительные строки допускаются"),
            ("Возмещаемые расходы", "объединенная ячейка по спецификации"),
            ("Невозмещаемые расходы", "объединенная ячейка по спецификации"),
            ("№ счф", "объединенная ячейка; все документы через Enter, без сокращений количества документов"),
            ("(+/-)", "объединенная ячейка по спецификации"),
        ],
        [2350, 7010],
        font_size=8,
    )
    add_body(
        doc,
        "Копирование в буфер должно формировать плоский TSV: значения объединенных ячеек повторяются в каждой строке, чтобы бухгалтер мог вставить данные в Excel без потери контекста.",
    )
    add_image(doc, IMG_XLSX, "Рисунок 4. Требуемый вид XLSX-выгрузки по свежему кейсу 660/1.")


def add_logging(doc):
    add_heading(doc, "13. Журнал расхождений", 1)
    add_body(
        doc,
        "Нужна таблица логирования не только для текущего экрана, но и для последующего анализа частоты проблем: какие типы операций чаще ломаются, где нет ключей 1С, какие поставщики/услуги дают расхождения.",
    )
    add_table(
        doc,
        ["Поле", "Назначение"],
        [
            ("run_id", "идентификатор запуска сверки"),
            ("contact_id, client_id, dog_id, spec_id", "контекст клиента, ЮЛ, договора и поставки"),
            ("doc_type, erp_doc_id, onec_doc_key, onec_guid", "идентификация документа с обеих сторон"),
            ("mismatch_code", "код статуса: AMOUNT_MISMATCH, VAT_MISMATCH, MISSING_1C и т.д."),
            ("erp_value, onec_value, diff_amount", "значения и числовая разница"),
            ("operation_id, operation_type_id, operation_type_name", "операция поставки и ее тип"),
            ("f_isvozm, source_table, source_row_id", "классификация и источник ERP-строки"),
            ("created_at, created_by", "время и пользователь"),
        ],
        [3000, 6360],
        font_size=8.1,
    )


def add_control_case(doc):
    add_heading(doc, "14. Контрольный кейс 660/1", 1)
    add_body(
        doc,
        "Проверка выполнялась на выбранных свежих поставках договора 660/1. Это не расчет по всему договору целиком. Итоговый расчет ниже воспроизводится из MariaDB ERP и должен совпадать с production-логикой ERP.",
    )
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ("Клиент", "АЭРО-ТРЕЙД"),
            ("contact_id / client_id / dog_id", "115 / 221 / 88"),
            ("Договор", "660/1, код 1С БП-003453"),
            ("Состав набора", "8 выбранных поставок договора 660/1, а не весь договор целиком."),
            ("Поставки", "spec_id 20351/1063, 20352/1064, 20353/1065, 20818/1068, 20824/1072, 20825/1073, 20822/1074, 20858/1076."),
            ("Итоги", "поступления/get_paidsum 2 331 111,33; реализация/get_realizsum: возмещаемые 2 152 392,27 и невозмещаемые 263 133,46; сальдо -84 414,40"),
            ("Особенность", "`get_paidsum` совпал с `SUM(veda_acchist_docs.f_clssum)`; подрядные расходы без клиентской реализации исключаются из взаиморасчета, но логируются"),
        ],
        [2550, 6810],
        font_size=8.2,
    )
    add_table(
        doc,
        ["spec_id", "Номер", "Тип", "Подтип", "Дата", "Наименование ERP"],
        [
            ("20351", "1063", "Заявка", "Заявка", "2025-05-12", "660/1/1063/ВА/АЭРО-ТРЕЙД"),
            ("20352", "1064", "Заявка", "Заявка", "2025-05-12", "660/1/1064/ВА/АЭРО-ТРЕЙД"),
            ("20353", "1065", "Заявка", "Заявка", "2025-05-12", "660/1/1065/ВА/АЭРО-ТРЕЙД"),
            ("20818", "1068", "Заявка", "Заявка", "2025-05-26", "660/1/1068/ВА/АЭРО-ТРЕЙД"),
            ("20824", "1072", "Заявка", "Заявка", "2025-05-26", "660/1/1072/ВА/АЭРО-ТРЕЙД"),
            ("20825", "1073", "Заявка", "Заявка", "2025-05-26", "660/1/1073/ВА/АЭРО-ТРЕЙД"),
            ("20822", "1074", "Заявка", "Заявка", "2025-05-26", "660/1/1074/ВА/АЭРО-ТРЕЙД"),
            ("20858", "1076", "Заявка", "Заявка", "2025-05-26", "660/1/1076/ВА/АЭРО-ТРЕЙД"),
        ],
        [850, 850, 1050, 1050, 1300, 4260],
        font_size=7.5,
    )
    add_table(
        doc,
        ["Спец.", "Оплата", "Возм.", "Невозм.", "Реализация", "(+/-)"],
        [
            ("1063", "199 375,04", "185 181,88", "41 275,82", "226 457,70", "-27 082,66"),
            ("1064", "134 234,42", "124 577,67", "27 882,63", "152 460,30", "-18 225,88"),
            ("1065", "170 517,04", "158 354,45", "35 211,68", "193 566,13", "-23 049,09"),
            ("1068", "442 296,92", "416 239,47", "36 640,04", "452 879,51", "-10 582,59"),
            ("1072", "319 193,45", "287 869,95", "26 049,03", "313 918,98", "5 274,47"),
            ("1073", "131 241,57", "118 363,63", "10 708,61", "129 072,24", "2 169,33"),
            ("1074", "442 296,92", "416 239,47", "36 640,04", "452 879,51", "-10 582,59"),
            ("1076", "491 955,97", "445 565,75", "48 725,61", "494 291,36", "-2 335,39"),
            ("Итого", "2 331 111,33", "2 152 392,27", "263 133,46", "2 415 525,73", "-84 414,40"),
        ],
        [1000, 1670, 1670, 1670, 1670, 1680],
        font_size=7.9,
        alignments=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
            WD_ALIGN_PARAGRAPH.RIGHT,
        ],
    )
    add_callout(
        doc,
        "Проверка источника",
        "Контрольный набор собран серверным SQL-скриптом из MariaDB ERP. Для каждой поставки проверено равенство `get_paidsum` и распределенной оплаты `SUM(veda_acchist_docs.f_clssum)`, а дельта рассчитана как `SUM(get_paidsum) - SUM(get_realizsum)`.",
    )
    add_callout(
        doc,
        "Что этот кейс не доказывает",
        "Контрольный кейс 660/1 подтверждает корректность ERP-сальдовки по выбранным поставкам. Он не подтверждает, что сальдо ERP и 1С совпадает по договору за период. Для такого вывода нужен отдельный запуск страницы `Монитор ERP ↔ 1С` с сальдо на начало, оборотами и сальдо на конец.",
        fill=WARN_FILL,
    )
    add_heading(doc, "14.1 Проверка монитора ERP ↔ 1С", 2)
    add_body(
        doc,
        "Для страницы монитора выполнена проверка по тому же договору 660/1 на небольшом онлайн-срезе. В рабочем контуре источник 1С должен предоставляться текущим сервисом 1С с теми же документами и реквизитами, которые перечислены в маппинге.",
    )
    add_table(
        doc,
        ["Параметр", "Результат"],
        [
            ("Параметры проверки", "ЮЛ `client_id=221`, договор `dog_id=88`, 3 свежие поставки договора 660/1"),
            ("Объем", "3 поставки, 27 строк документной сверки"),
            ("Итог", "26 строк `ОК`, 1 строка с расхождением договора/заявки"),
            ("Пример расхождения", "`spec_id=20825`, Заявка №1073, ERP-операция `Портовые расходы. Поручение на оплату`, сумма 1 456,59; 1С документ `00БП-013353` найден по дате/сумме, но аналитика договора/заявки не совпала"),
            ("Вывод", "Такой результат нельзя показывать как общий `ОК`; нужен отдельный статус `Договор/заявка расходится` и запись в журнал расхождений."),
            ("Производительность", "Онлайн-срез из 3 поставок занимает десятки секунд; 10 поставок не подходят для синхронного экрана. Полный договор/историю считать фоновым обработчиком."),
        ],
        [2300, 7060],
        font_size=7.65,
    )
    add_body(
        doc,
        "Письма и внешние корректировки вида `по письму` в ERP отсутствуют и не должны искусственно добавляться в расчет. Если источник отсутствует, строка остается пустой или получает статус `Недостаточно данных`.",
        color=RISK,
    )


def add_acceptance(doc):
    add_heading(doc, "15. Критерии приемки", 1)
    checks = [
        "Текущий экран называется и описывается как `Сальдо по поставкам`, а не как полноценная сверка ERP ↔ 1С.",
        "Экран сальдо строит иерархию `veda_contacts -> veda_clients -> veda_dogs -> veda_specs` и не смешивает уровни цифрами.",
        "Свежий кейс 660/1 воспроизводит итоги ERP: поступления/get_paidsum 2 331 111,33; реализация/get_realizsum 2 415 525,73; сальдо -84 414,40. Это критерий ERP-сальдовки, не критерий совпадения с 1С.",
        "В названии строки поставки используется тип из исходных таблиц ERP (`veda_specs.f_typez/f_subtype` + `veda_spr`), например `Заявка №1063`, а не хардкод `Спецификация {num}` и не view.",
        "На онлайн-экране действует малый лимит и пагинация: срез из 3 поставок по ЮЛ `client_id=221` и договору `dog_id=88` отрабатывает, а попытка 10 поставок уже неприемлема для синхронного запроса.",
        "Сумма оплаты берется из `veda_acchist_docs.f_clssum`; полная сумма банка не используется как оплата поставки.",
        "`f_isvozm=1` идет в возмещаемые расходы; `f_isvozm=2` идет в невозмещаемые расходы через `get_realizsum`; подрядные расходы без клиентской реализации исключаются из сальдо.",
        "`Нет СФ/УПД` определяется отсутствием закрывающего документа после загрузки источников, а `Вопрос по НДС` - только несовпадением ставки НДС ERP и 1С.",
        "На странице сальдо не показываются счетчики совпавших документов 1С; документное сопоставление находится только в `Мониторе ERP ↔ 1С`.",
        "Для настоящей сверки ERP ↔ 1С есть отдельная страница `Монитор ERP ↔ 1С` с запуском `run_id`, статусом источников, блоком `ERP без 1С`, блоком `1С без ERP`, блоком расхождений по документам и блоком сальдо.",
        "Отдельная страница сверки показывает начальное сальдо, обороты и конечное сальдо ERP/1С; при произвольном наборе поставок предупреждает, что полноценный акт сверки невозможен.",
        "Фоновый обработчик для тяжелой исторической сверки сохраняет результат и не блокирует интерфейс синхронным пересчетом всей истории.",
        "XLSX выгружает все документы в объединенной ячейке через Enter, без сокращения списка до счетчика количества документов.",
        "Журнал расхождений пишет тип ошибки, источник, операцию, суммы и контекст поставки.",
    ]
    for check in checks:
        add_bullet(doc, check)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = ""
    run = p.add_run("ТЗ: сальдо по поставкам + сверка ERP ↔ 1С")
    set_font(run, size=8, color=MUTED)


def build_doc():
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    title_page(doc)
    doc.add_page_break()
    add_user_requirements(doc)
    add_product_split(doc)
    add_workflow(doc)
    add_functional_requirements(doc)
    add_reconciliation_page_requirements(doc)
    add_runner_requirements(doc)
    add_nonfunctional_requirements(doc)
    add_hierarchy(doc)
    add_ui(doc)
    add_mapping(doc)
    add_field_dictionary(doc)
    add_sources_and_selects(doc)
    add_column_rules(doc)
    add_statuses(doc)
    add_excel(doc)
    add_logging(doc)
    add_control_case(doc)
    add_acceptance(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    shutil.copyfile(OUT, OUT_VERSIONED)
    print(OUT)
    print(OUT_VERSIONED)


if __name__ == "__main__":
    build_doc()
